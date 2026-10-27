# /// script
# requires-python = ">=3.11"
# dependencies = ["python-pptx", "python-docx", "pymupdf"]
# ///
"""Create a draft brand from a .pptx (often a client's or corporate deck).

    rep = brand_from_pptx(Path("deck.pptx"), "acme", Path("brands"))
    print(format_report(rep))      # texts and images a human must review
    confirm_brand(rep.brand_dir)   # after review: the brand can be loaded

CLI: `uv run brand_from_pptx.py SRC NAME [--out DIR] [--no-samples]` prints the report;
`uv run brand_from_pptx.py --confirm DIR` confirms a reviewed draft. Default `--out` is
~/.mle-setup/brands.

The brand folder gets:
- `template.pptx`: SRC with no slides, notes, comments, comment authors, edit history,
  printer settings, thumbnail, customXml or non-label custom properties; core properties and the app.xml
  Company/Manager cleared, slide titles removed from app.xml. A sensitivity label
  (`docMetadata/LabelInfo.xml`, or MSIP_Label_* custom properties) is kept untouched.
- `assets-review/`: every image used by masters/layouts, and the rendered samples.
- `brand.toml` with `draft = true` (loading refuses it until `--confirm`), style "outlined":
  theme dk2 -> primary, accent1 -> accent, accent2 -> highlight, dk1 -> text, lt1 -> background,
  lt2 -> surface; muted = text mixed 45% toward background (a mid grey that keeps ~4.5:1 on a
  white background); on_primary = lt1 or dk1, whichever contrasts more with primary;
  card_border = accent (outlined cards are drawn in the accent); done/planned come from the
  preset. Fonts: theme major latin -> heading, minor latin -> body. Layout roles are guessed
  from layout names.
The report lists every non-empty text in masters and layouts (placeholder prompts included) so
a human can spot client data left in the template before confirming.
"""

from __future__ import annotations

import argparse
import re
import shutil
import sys
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.exc import PackageNotFoundError
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from brand import COLOR_TOKENS, PLUGIN_BRANDS, BrandError, load_brand
from new_brand import BLACK, DEFAULT_OUT, WHITE, best_contrast, brand_folder, dump_toml, mix, sample_slides

STYLE = "outlined"
MUTED_MIX = 0.45
EMU_PER_IN = 914400
LABEL_PART = "docMetadata/LabelInfo.xml"
LABEL_PROP = "MSIP_Label_"  # legacy sensitivity label custom properties
THEME_MAP = {"primary": "dk2", "accent": "accent1", "highlight": "accent2",
             "text": "dk1", "background": "lt1", "surface": "lt2"}
SYS_COLORS = {"windowText": BLACK, "window": WHITE}  # sysClr without lastClr
LAYOUT_GUESSES = {
    "title": ("title slide",),
    "content": ("title and content", "title & content", "title only"),
    "section": ("section",),
}
CORE_FIELDS = ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category",
               "content_status", "identifier")
# Presentation-level parts with people, comments, edit history, document or device metadata (not design).
DROP_PRES_RELS = ("/commentAuthors", "/authors", "/revisionInfo", "/changesInfo", "/customXml",
                  "/printerSettings")
APP_CLEAR = ("Company", "Manager", "HyperlinkBase")
APP_DROP = ("HeadingPairs", "TitlesOfParts", "HLinks")  # TitlesOfParts lists slide titles
APP_ZERO = ("Slides", "Notes", "HiddenSlides", "MMClips", "Words", "Paragraphs")
LEFTOVERS = re.compile(r"ppt/(slides|notesSlides|comments|changesInfos)/|ppt/(commentAuthors|authors|"
                       r"revisionInfo)\.xml|ppt/printerSettings/|docProps/thumbnail|customXml/")
P14_SECTIONS = "{http://schemas.microsoft.com/office/powerpoint/2010/main}sectionLst"
DRAFT_NOTE = "DRAFT:"


@dataclass
class ExtractReport:
    brand_dir: Path
    layouts: list[str]
    texts: list[tuple[str, str]]  # (location like "layout 'Title Slide' / TextBox 3", text)
    images: list[Path]
    has_label: bool
    samples: list[Path] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)  # mapping fallbacks and warnings


# ── reading ──
def _related(part, reltype: str):
    try:
        return part.part_related_by(reltype)
    except KeyError:
        return None


def _xml(part):
    return part._element if hasattr(part, "_element") else etree.fromstring(part.blob)


def _design_parts(prs) -> list[tuple[str, str, object]]:
    """(location, file prefix, part) for every master and layout, then the notes/handout masters."""
    out, n = [], 0
    for i, master in enumerate(prs.slide_masters, 1):
        out.append((f"master '{master.name}'" if master.name else f"master {i}", f"master{i}", master.part))
        for layout in master.slide_layouts:
            n += 1
            out.append((f"layout '{layout.name}'", f"layout{n:02d}", layout.part))
    for reltype, where, prefix in ((RT.NOTES_MASTER, "notes master", "notesmaster"),
                                   (RT.HANDOUT_MASTER, "handout master", "handoutmaster")):
        part = _related(prs.part, reltype)
        if part is not None:
            out.append((where, prefix, part))
    return out


def _texts(root, where: str) -> list[tuple[str, str]]:
    out = []
    for el in root.iter(qn("p:sp"), qn("p:graphicFrame")):
        cnv = el.find("./*/" + qn("p:cNvPr"))
        shape = cnv.get("name", "") if cnv is not None else ""
        paras = ("".join(t.text or "" for t in p.iter(qn("a:t"))) for p in el.iter(qn("a:p")))
        text = "\n".join(x for x in paras if x.strip()).strip()
        if text:
            out.append((f"{where} / {shape}", text))
    return out


def _copy_images(parts, dest: Path) -> list[Path]:
    copied: dict[str, Path] = {}
    for _, prefix, part in parts:
        for rel in part.rels.values():
            if rel.is_external or rel.reltype != RT.IMAGE or rel.target_part.partname in copied:
                continue
            img = rel.target_part
            dest.mkdir(parents=True, exist_ok=True)
            path = dest / f"{prefix}-{Path(img.partname).name}"
            path.write_bytes(img.blob)
            copied[img.partname] = path
    return list(copied.values())


def _theme(master) -> tuple[dict[str, str], str, str]:
    """({dk1: hex, ...}, major latin typeface, minor latin typeface) of the master's theme."""
    part = _related(master.part, RT.THEME)
    if part is None:
        return {}, "", ""
    root = etree.fromstring(part.blob)
    colors = {}
    scheme = root.find(".//" + qn("a:clrScheme"))
    for slot in scheme if scheme is not None else []:
        if not isinstance(slot.tag, str) or not len(slot):
            continue
        c = slot[0]
        if c.tag == qn("a:srgbClr"):
            value = c.get("val")
        elif c.tag == qn("a:sysClr"):
            value = c.get("lastClr") or SYS_COLORS.get(c.get("val", ""))
        else:
            value = None
        if value and re.fullmatch(r"[0-9A-Fa-f]{6}", value):
            colors[etree.QName(slot).localname] = value.upper()

    def font(kind: str) -> str:
        latin = root.find(f".//{qn('a:fontScheme')}/{qn(kind)}/{qn('a:latin')}")
        return (latin.get("typeface") or "").strip() if latin is not None else ""

    return colors, font("a:majorFont"), font("a:minorFont")


def _colors(scheme: dict[str, str], notes: list[str]) -> dict[str, str]:
    colors = {tok: scheme[slot] for tok, slot in THEME_MAP.items() if slot in scheme}
    missing = [f"{tok} ({slot})" for tok, slot in THEME_MAP.items() if slot not in scheme]
    if missing:
        notes.append(f"theme colors missing, preset defaults used: {', '.join(missing)}")
    if "text" in colors and "background" in colors:
        colors["muted"] = mix(colors["text"], colors["background"], MUTED_MIX)
    if "primary" in colors:
        colors["on_primary"] = best_contrast(colors["primary"], scheme.get("lt1", WHITE), scheme.get("dk1", BLACK))
    if "accent" in colors:
        colors["card_border"] = colors["accent"]
    return {k: colors[k] for k in COLOR_TOKENS if k in colors}


def _fonts(major: str, minor: str, notes: list[str]) -> dict[str, str]:
    def usable(f: str) -> bool:
        return bool(f) and not f.startswith("+")  # "+mj-lt" etc. point back at the theme

    heading, body = (f if usable(f) else "" for f in (major, minor))
    if not (heading or body):
        fallback = load_brand(PLUGIN_BRANDS / "neutral").fonts["body"]
        notes.append(f"theme fonts unusable ({major!r}, {minor!r}): using the neutral brand's {fallback!r}")
        heading = body = fallback
    elif not heading:
        notes.append(f"theme major font unusable ({major!r}): heading uses the body font {body!r}")
        heading = body
    elif not body:
        notes.append(f"theme minor font unusable ({minor!r}): body uses the heading font {heading!r}")
        body = heading
    return {"heading": heading, "body": body}


def _guess_layouts(names: list[str]) -> dict[str, str]:
    out = {}
    for role, patterns in LAYOUT_GUESSES.items():
        hit = next((n for p in patterns for n in names if p in n.lower()), None)
        if hit:
            out[role] = hit
    return out


# ── cleaning ──
def _filtered_xml(blob: bytes, edit) -> bytes:
    root = etree.fromstring(blob)
    edit(root)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _clean_app(root) -> None:
    for el in list(root):
        local = etree.QName(el).localname
        if local in APP_CLEAR:
            el.text = None
        elif local in APP_DROP:
            root.remove(el)
        elif local in APP_ZERO:
            el.text = "0"


def _remove_all(root, tag: str) -> None:
    for el in list(root.iter(tag)):
        el.getparent().remove(el)


def _strip(prs) -> None:
    """Remove slides and everything that is about the document rather than the design."""
    pres = prs.part
    root = pres._element
    _remove_all(root, qn("p:custShowLst"))  # custom shows list slides
    for ext in list(root.iter(qn("p:ext"))):
        if ext.find(P14_SECTIONS) is not None:  # section names and their slide ids
            ext.getparent().remove(ext)
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        ids.remove(sld)
        pres.drop_rel(sld.rId)
    for rId, rel in list(pres.rels.items()):
        if rel.reltype.endswith(DROP_PRES_RELS):
            pres.drop_rel(rId)

    props = prs.core_properties
    for name in CORE_FIELDS:
        setattr(props, name, "")

    pkg = pres.package
    for rId, rel in list(pkg._rels.items()):
        if rel.is_external:
            continue
        if rel.reltype == RT.THUMBNAIL:  # a picture of slide 1
            pkg.drop_rel(rId)
        elif rel.reltype == RT.EXTENDED_PROPERTIES:
            rel.target_part.blob = _filtered_xml(rel.target_part.blob, _clean_app)
        elif rel.reltype == RT.CUSTOM_PROPERTIES:
            root_cp = etree.fromstring(rel.target_part.blob)
            for prop in list(root_cp):
                if not (prop.get("name") or "").startswith(LABEL_PROP):
                    root_cp.remove(prop)
            if len(root_cp):
                rel.target_part.blob = etree.tostring(root_cp, xml_declaration=True, encoding="UTF-8",
                                                      standalone=True)
            else:
                pkg.drop_rel(rId)

    # Any other way back to slides (e.g. viewProps' outline view) goes too.
    for part in list(pkg.iter_parts()):
        for rId, rel in list(part.rels.items()):
            if not rel.is_external and rel.target_part.content_type in (CT.PML_SLIDE, CT.PML_NOTES_SLIDE):
                part.rels.pop(rId)
        if part.content_type == CT.PML_VIEW_PROPS:
            part.blob = _filtered_xml(part.blob, lambda r: _remove_all(r, qn("p:sldLst")))
        elif part.content_type == CT.PML_PRES_PROPS:
            part.blob = _filtered_xml(part.blob, lambda r: _remove_all(r, qn("p:custShow")))


def _verify(template: Path, src: Path) -> bool:
    """Reopen the template and check the package; return whether a sensitivity label is kept."""
    if len(Presentation(str(template)).slides):
        raise BrandError(f"{template}: slides left in the template")
    with zipfile.ZipFile(template) as z:
        names = set(z.namelist())
        overrides = re.findall(r'PartName="/([^"]+)"', z.read("[Content_Types].xml").decode("utf-8"))
        label = z.read(LABEL_PART) if LABEL_PART in names else None
        custom = z.read("docProps/custom.xml") if "docProps/custom.xml" in names else b""
    leftovers = sorted(n for n in names if LEFTOVERS.match(n))
    missing = [o for o in overrides if o not in names]
    if leftovers or missing:
        raise BrandError(f"{template}: invalid template (leftover parts {leftovers}, missing parts {missing})")
    with zipfile.ZipFile(src) as z:
        src_label = z.read(LABEL_PART) if LABEL_PART in z.namelist() else None
    if src_label is not None and label != src_label:
        raise BrandError(f"{template}: the sensitivity label ({LABEL_PART}) was not kept intact")
    return label is not None or LABEL_PROP.encode() in custom


# ── samples ──
def _render_samples(rep: ExtractReport) -> list[Path]:
    """One slide per layout plus the standard sample slides, rendered to PNG. Never fails."""
    try:
        from deck import Deck
        from render import render

        d = Deck(load_brand(rep.brand_dir, allow_draft=True))
        for layout in d.prs.slide_layouts:
            s = d.prs.slides.add_slide(layout)
            if s.shapes.title is not None:
                s.shapes.title.text = layout.name
        sample_slides(d)
        out = rep.brand_dir / "assets-review" / "samples" / "sample.pptx"
        out.parent.mkdir(parents=True, exist_ok=True)
        d.save(out)
        return render(out, out_dir=out.parent).pngs
    except Exception as e:  # samples help the review but are optional (no Office, render timeout, ...)
        msg = f"samples not rendered: {type(e).__name__}: {e}"
        print(msg, file=sys.stderr)
        rep.notes.append(msg)
        return []


# ── public API ──
def brand_from_pptx(src: Path, name: str, out_dir: Path, render_samples: bool = True) -> ExtractReport:
    """Write the draft brand `out_dir/name` from `src`; return what a human must review."""
    src = Path(src).expanduser()
    if not src.is_file():
        raise BrandError(f"{src}: file not found")
    folder = brand_folder(out_dir, name)
    prs = Presentation(str(src))
    notes: list[str] = []
    parts = _design_parts(prs)
    layouts = list(dict.fromkeys(lay.name for m in prs.slide_masters for lay in m.slide_layouts))
    if len(prs.slide_masters) > 1:
        notes.append(f"{len(prs.slide_masters)} slide masters: decks use the layouts of the first one")
    texts = [t for where, _, part in parts for t in _texts(_xml(part), where)]
    scheme, major, minor = _theme(prs.slide_masters[0])
    colors = _colors(scheme, notes)
    fonts = _fonts(major, minor, notes)
    guesses = _guess_layouts([lay.name for lay in prs.slide_layouts])
    notes.append(f"slide size {prs.slide_width / EMU_PER_IN:.2f} x {prs.slide_height / EMU_PER_IN:.2f} in")
    unmapped = [r for r in LAYOUT_GUESSES if r not in guesses]
    if unmapped:
        notes.append(f"no layout guessed for {', '.join(unmapped)}: those slides are drawn, or map them in "
                     "[deck.layouts]")

    created = not folder.exists()
    try:
        images = _copy_images(parts, folder / "assets-review")
        _strip(prs)
        template = folder / "template.pptx"
        folder.mkdir(parents=True, exist_ok=True)
        prs.save(str(template))
        has_label = _verify(template, src)
        data = {"name": name, "description": "Extracted from a .pptx template", "draft": True, "style": STYLE,
                "colors": colors, "fonts": fonts,
                "deck": {"template": template.name, "footer": "", "layouts": guesses}}
        header = "\n".join([
            f"{DRAFT_NOTE} written by brand_from_pptx.py. Review the reported texts, assets-review/ and",
            f"{DRAFT_NOTE} the samples for client data, then run: brand_from_pptx.py --confirm <this folder>",
            "Theme mapping: dk2 primary, accent1 accent, accent2 highlight, dk1 text, lt1 background,",
            f"lt2 surface; muted = text {MUTED_MIX:.0%} toward background; on_primary = lt1 or dk1 (best",
            f"contrast); card_border = accent; other tokens from the {STYLE!r} preset.",
            *notes,
        ])
        (folder / "brand.toml").write_text(dump_toml(data, header), encoding="utf-8")
    except BaseException:
        if created:
            shutil.rmtree(folder, ignore_errors=True)
        raise
    rep = ExtractReport(folder, layouts, texts, images, has_label, notes=notes)
    if render_samples:
        rep.samples = _render_samples(rep)
    return rep


def confirm_brand(brand_dir: Path) -> None:
    """Remove the draft flag (and the draft comments) from brand.toml, then check it loads."""
    path = Path(brand_dir).expanduser() / "brand.toml"
    text = path.read_text(encoding="utf-8")
    m = re.search(r"(?m)^[ \t]*\[", text)
    head, rest = (text[:m.start()], text[m.start():]) if m else (text, "")
    head = re.sub(rf"(?m)^[ \t]*(draft[ \t]*=.*|# {DRAFT_NOTE}.*)(\r?\n|$)", "", head)
    path.write_text(head + rest, encoding="utf-8")
    load_brand(path.parent)


def format_report(rep: ExtractReport) -> str:
    lines = [f"draft brand: {rep.brand_dir}",
             f"sensitivity label: {'kept' if rep.has_label else 'none'}",
             f"layouts ({len(rep.layouts)}): {', '.join(rep.layouts)}",
             f"texts in masters/layouts ({len(rep.texts)}) - check for client data:"]
    lines += [f"  {where}: {text!r}" for where, text in rep.texts] or ["  (none)"]
    lines.append(f"images used by masters/layouts ({len(rep.images)}):")
    lines += [f"  {p}" for p in rep.images] or ["  (none)"]
    lines.append(f"samples ({len(rep.samples)}):")
    lines += [f"  {p}" for p in rep.samples] or ["  (none)"]
    lines += [f"note: {n}" for n in rep.notes]
    lines.append(f"next: review brand.toml and the above, then: brand_from_pptx.py --confirm {rep.brand_dir}")
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="brand_from_pptx.py", description="Create a draft brand from a .pptx.")
    ap.add_argument("src", nargs="?", help="the .pptx to extract the design from")
    ap.add_argument("name", nargs="?", help="brand (folder) name")
    ap.add_argument("--out", default=str(DEFAULT_OUT), metavar="DIR", help=f"default: {DEFAULT_OUT}")
    ap.add_argument("--no-samples", action="store_true", help="do not render sample slides")
    ap.add_argument("--confirm", metavar="DIR", help="confirm a reviewed draft brand folder")
    args = ap.parse_args(argv)
    if args.confirm and (args.src or args.name):
        ap.error("--confirm DIR takes no SRC/NAME")
    if not args.confirm and not (args.src and args.name):
        ap.error("SRC and NAME are required (or --confirm DIR)")
    try:
        if args.confirm:
            confirm_brand(Path(args.confirm))
            print(f"confirmed: {Path(args.confirm).expanduser()}")
        else:
            rep = brand_from_pptx(Path(args.src), args.name, Path(args.out), render_samples=not args.no_samples)
            print(format_report(rep))
    except (BrandError, OSError, ValueError, PackageNotFoundError, zipfile.BadZipFile) as e:
        print(f"brand_from_pptx failed: {e}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
