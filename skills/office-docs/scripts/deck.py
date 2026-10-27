"""Deck: an editable .pptx built from a brand folder.

    d = Deck("neutral")
    s = d.blank_slide()
    x, w = d.col(0, 6)  # first half of the grid
    d.text(s, x, 1.0, w, 0.8, "Hola", size=28, font="heading")
    d.save("out.pptx")

Public units are inches (floats); colors are a brand token ("accent") or a hex.
With `[deck] template` the deck starts from that .pptx (its slides removed) and the slide
factories use the layouts mapped in `[deck.layouts]`.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.slide import Slide, SlideLayout
from pptx.util import Emu

from brand import COLOR_TOKENS, Brand, BrandError, load_brand, normalize_hex
from deck_components import Components
from deck_prims import EMU_PER_IN, Primitives

SLIDE_W_EMU = 12192000  # 13.333 in (16:9)
SLIDE_H_EMU = 6858000  # 7.5 in
BLANK_LAYOUT = 6  # "Blank" in python-pptx's default template
FONT_ROLES = ("body", "heading", "mono")


class Deck(Primitives, Components):
    def __init__(self, brand: str | Path | Brand, footer: str | None = None, cwd: Path | None = None):
        self.brand = brand if isinstance(brand, Brand) else load_brand(brand, cwd=cwd)
        self.footer_text = self.brand.footer if footer is None else footer
        template = self.brand.template
        if template is None:
            self.prs = Presentation()
            self.prs.slide_width = Emu(SLIDE_W_EMU)
            self.prs.slide_height = Emu(SLIDE_H_EMU)
        else:
            if not Path(template).is_file():
                raise BrandError(f"deck.template: file not found: {template}")
            self.prs = Presentation(str(template))  # in memory: the template file is never written
            self._drop_slides()
        self.W = self.prs.slide_width / EMU_PER_IN
        self.H = self.prs.slide_height / EMU_PER_IN
        self.page = 0

    def color(self, c: str) -> RGBColor:
        """A brand token ('accent') or a hex ('2870CF', '#2870cf')."""
        if isinstance(c, RGBColor):
            return c
        if isinstance(c, str) and c in self.brand.colors:
            return RGBColor.from_string(self.brand.colors[c])
        try:
            return RGBColor.from_string(normalize_hex(c, "color"))
        except BrandError:
            raise ValueError(
                f"color {c!r} is neither a token nor a 6-digit hex (tokens: {', '.join(COLOR_TOKENS)})"
            ) from None

    def font(self, role: str) -> str:
        """'body' | 'heading' | 'mono' map to the brand fonts; any other string is a literal font name."""
        return self.brand.fonts[role] if role in FONT_ROLES else role

    def col(self, start: int, span: int = 1) -> tuple[float, float]:
        """x and width (inches) of grid columns [start, start+span), 0-based."""
        g = self.brand.grid
        n = int(g["columns"])
        if start < 0 or span < 1 or start + span > n:
            raise ValueError(f"col({start}, {span}) outside a {n}-column grid (start is 0-based)")
        col_w = (self.W - 2 * g["margin"] - (n - 1) * g["gutter"]) / n
        return g["margin"] + start * (col_w + g["gutter"]), span * col_w + (span - 1) * g["gutter"]

    def _drop_slides(self) -> None:
        """Remove every slide (sldId entry and its relationship) so no slide part is saved."""
        ids = self.prs.slides._sldIdLst
        for sld in list(ids):
            ids.remove(sld)
            self.prs.part.drop_rel(sld.rId)

    def layout(self, role: str) -> SlideLayout | None:
        """Template layout mapped to `role` in [deck.layouts]; None without a template or mapping."""
        name = self.brand.layouts.get(role)
        if self.brand.template is None or name is None:
            return None
        layouts = self.prs.slide_layouts
        found = next((l for l in layouts if l.name == name), None)
        if found is None:
            raise BrandError(
                f"deck.layouts.{role}: layout {name!r} not found in {self.brand.template}; "
                f"available: {[l.name for l in layouts]}"
            )
        return found

    def _blank_layout(self) -> SlideLayout:
        layouts = self.prs.slide_layouts
        if self.brand.template is None:
            return layouts[BLANK_LAYOUT]
        named = next((l for l in layouts if l.name == "Blank"), None)
        return named or min(layouts, key=lambda l: len(l.placeholders))

    def blank_slide(self, bg: str | None = None) -> Slide:
        """New empty slide; `bg=None` uses components.content_bg.

        With a template: its "Blank" layout (else the one with fewest placeholders), placeholders removed.
        """
        s = self.prs.slides.add_slide(self._blank_layout())
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = self.color(bg or self.brand.components["content_bg"])
        self.page += 1
        return s

    def notes(self, slide: Slide, text: str) -> Slide:
        slide.notes_slide.notes_text_frame.text = text.strip()
        return slide

    def save(self, path: str | Path) -> int:
        self.prs.save(str(path))
        return len(self.prs.slides)
