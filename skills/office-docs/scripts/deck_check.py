"""Advisory layout check for a built deck: heuristics that flag likely layout problems.

    issues = check_file("out.pptx")   # or check_presentation(deck.prs)

Kinds: off_slide (bbox past the slide edge), overflow (text taller than its box, autofit off),
overlap (two text shapes crossing, neither containing the other), covered (a filled shape or a
picture drawn after a text shape crosses it without containing it), long_title (> 55 chars).
It only reports; it never raises on issues and never blocks a build.
"""

from __future__ import annotations

from dataclasses import dataclass
from itertools import combinations
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn

from deck_prims import EMU_PER_IN, est_lines

OFF_SLIDE_TOL = 0.02  # inches
OVERFLOW_FACTOR = 1.05
LINE_HEIGHT = 1.2  # x font size
DEFAULT_SIZE_PT = 18.0
DEFAULT_INSETS = {"lIns": 0.1, "rIns": 0.1, "tIns": 0.05, "bIns": 0.05}  # PowerPoint defaults, inches
TITLE_MAX_CHARS = 55
TITLE_PLACEHOLDERS = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
FILLS = tuple(qn(f"a:{k}") for k in ("solidFill", "gradFill", "pattFill", "blipFill", "grpFill"))
COVER_TOL = 0.01  # inches: overlaps thinner than this (rounding, touching edges) do not count


@dataclass
class Issue:
    slide: int  # 1-based
    shape: str
    kind: str  # off_slide | overflow | overlap | covered | long_title
    detail: str

    def __str__(self) -> str:
        return f"slide {self.slide}: {self.kind} {self.shape} — {self.detail}"


def _bbox(sh) -> tuple[float, float, float, float] | None:
    """(left, top, right, bottom) in inches; None when the shape inherits its geometry."""
    if None in (sh.left, sh.top, sh.width, sh.height):
        return None
    l, t = sh.left / EMU_PER_IN, sh.top / EMU_PER_IN
    return l, t, l + sh.width / EMU_PER_IN, t + sh.height / EMU_PER_IN


def _text(sh) -> str:
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


def _off_slide(box, W, H) -> bool:
    l, t, r, b = box
    return l < -OFF_SLIDE_TOL or t < -OFF_SLIDE_TOL or r > W + OFF_SLIDE_TOL or b > H + OFF_SLIDE_TOL


def _overflow(sh, box) -> str | None:
    """Detail when the estimated text height exceeds 1.05 x the box height minus its insets."""
    body = sh.text_frame._txBody.find(qn("a:bodyPr"))
    if body is not None and (body.find(qn("a:normAutofit")) is not None
                             or body.find(qn("a:spAutoFit")) is not None):
        return None
    ins = {k: (int(body.get(k)) / EMU_PER_IN if body is not None and body.get(k) is not None else v)
           for k, v in DEFAULT_INSETS.items()}
    width = box[2] - box[0] - ins["lIns"] - ins["rIns"]
    avail = (box[3] - box[1] - ins["tIns"] - ins["bIns"]) * 72
    wrap = body is None or body.get("wrap") != "none"
    need = 0.0
    for p in sh.text_frame.paragraphs:
        sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
        size = max(sizes) if sizes else DEFAULT_SIZE_PT
        segments = p.text.split("\v")  # line breaks (<a:br>) start a new line
        lines = sum(est_lines(seg, size, width) for seg in segments) if wrap else len(segments)
        need += lines * size * LINE_HEIGHT
    if need <= OVERFLOW_FACTOR * avail:
        return None
    return f"text needs ~{need / 72:.2f} in, box fits {max(avail, 0) / 72:.2f} in"


def _is_title(sh) -> bool:
    return sh.name == "title" or (sh.is_placeholder and sh.placeholder_format.type in TITLE_PLACEHOLDERS)


def _opaque(sh) -> bool:
    """A picture, or an autoshape with a fill (explicit, or from its style when spPr sets none)."""
    el = sh._element
    if el.tag == qn("p:pic"):
        return True
    if el.tag != qn("p:sp"):
        return False
    sppr = el.find(qn("p:spPr"))
    for child in sppr if sppr is not None else ():
        if child.tag == qn("a:noFill"):
            return False
        if child.tag in FILLS:
            return True
    ref = el.find(f"{qn('p:style')}/{qn('a:fillRef')}")
    return ref is not None and ref.get("idx", "0") != "0"


def _crosses(a, b) -> bool:
    """Boxes intersect (positive area) and neither contains the other."""
    if not (a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]):
        return False

    def inside(x, y):
        return x[0] >= y[0] and x[1] >= y[1] and x[2] <= y[2] and x[3] <= y[3]

    return not (inside(a, b) or inside(b, a))


def _covers(s, t) -> bool:
    """Box `s` hides part of box `t`: they overlap by more than COVER_TOL and `s` does not contain `t`."""
    if min(s[2], t[2]) - max(s[0], t[0]) <= COVER_TOL or min(s[3], t[3]) - max(s[1], t[1]) <= COVER_TOL:
        return False
    tol = COVER_TOL
    return not (t[0] >= s[0] - tol and t[1] >= s[1] - tol and t[2] <= s[2] + tol and t[3] <= s[3] + tol)


def check_presentation(prs) -> list[Issue]:
    W, H = prs.slide_width / EMU_PER_IN, prs.slide_height / EMU_PER_IN
    issues: list[Issue] = []
    for n, slide in enumerate(prs.slides, 1):
        texts, drawn = [], []  # drawn: (name, box, has text, opaque) in z-order
        for sh in slide.shapes:
            box = _bbox(sh)
            text = _text(sh)
            if box is not None:
                if _off_slide(box, W, H):
                    issues.append(Issue(n, sh.name, "off_slide",
                                        f"bbox ({box[0]:.2f}, {box[1]:.2f})-({box[2]:.2f}, {box[3]:.2f}) in "
                                        f"exceeds the {W:.2f} x {H:.2f} in slide"))
                if text and (detail := _overflow(sh, box)):
                    issues.append(Issue(n, sh.name, "overflow", detail))
                if text:
                    texts.append((sh.name, box))
                drawn.append((sh.name, box, bool(text), _opaque(sh)))
            if text and _is_title(sh) and len(text) > TITLE_MAX_CHARS:
                issues.append(Issue(n, sh.name, "long_title", f"{len(text)} chars (max {TITLE_MAX_CHARS})"))
        for (na, a), (nb, b) in combinations(texts, 2):
            if _crosses(a, b):
                issues.append(Issue(n, na, "overlap", f"text crosses {nb}"))
        for i, (nt, t, has_text, _) in enumerate(drawn):
            if not has_text:
                continue
            for ns, s, s_text, opaque in drawn[i + 1:]:
                # two crossing text shapes are already an overlap
                if opaque and _covers(s, t) and not (s_text and _crosses(s, t)):
                    issues.append(Issue(n, nt, "covered", f"covered by {ns}, drawn later"))
    return issues


def check_file(path: str | Path) -> list[Issue]:
    return check_presentation(Presentation(str(path)))
