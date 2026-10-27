"""Primitives mixin for Deck: free-position shapes, text, images and effects.

Every method takes the slide first, positions in inches, colors as a brand token
or hex (resolved by `self.color`), fonts as a role or literal name (`self.font`),
and returns the created shape.
"""

from __future__ import annotations

from pathlib import Path

from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.parts.image import Image as PptxImage
from pptx.util import Emu, Pt

EMU_PER_IN = 914400
CAPTION_H = 0.45
CAPTION_GAP = 0.1

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
ARROWS = {
    "right": MSO_SHAPE.RIGHT_ARROW, "left": MSO_SHAPE.LEFT_ARROW,
    "up": MSO_SHAPE.UP_ARROW, "down": MSO_SHAPE.DOWN_ARROW,
}
RUN_KEYS = {"text", "size", "color", "bold", "italic", "font"}
FIT_MODES = ("contain", "cover", "stretch")


def emu(inches: float) -> Emu:
    return Emu(round(inches * EMU_PER_IN))


def est_lines(text: str, size_pt: float, width_in: float) -> int:
    """Lines the text will take, assuming an average glyph width of ~0.5 em."""
    cpl = max(10, int(width_in * 72 / (size_pt * 0.5)))
    return max(1, -(-len(text) // cpl))


def _pick(name: str, value: str, options: dict):
    if value not in options:
        raise ValueError(f"{name} {value!r} is not one of {', '.join(options)}")
    return options[value]


def _alpha_val(opacity: float, name: str = "opacity") -> str:
    if not 0 <= opacity <= 100:
        raise ValueError(f"{name} must be 0-100, got {opacity}")
    return str(int(opacity * 1000))


class Primitives:
    # ── shapes and lines ──
    def rect(self, s, x, y, w, h, fill="surface", line=None, line_w=1.0, dash=False, radius=None, opacity=None):
        """Rectangle; `radius` (0-0.5, fraction of the short side) makes it rounded; `fill=None` is transparent."""
        kind = MSO_SHAPE.RECTANGLE if radius is None else MSO_SHAPE.ROUNDED_RECTANGLE
        sh = s.shapes.add_shape(kind, emu(x), emu(y), emu(w), emu(h))
        sh.shadow.inherit = False
        if radius is not None:
            sh.adjustments[0] = radius
        self._fill(sh, fill)
        self._line(sh, line, line_w, dash)
        if opacity is not None:
            if fill is None:
                raise ValueError("opacity needs a fill color")
            srgb = sh._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            srgb.append(srgb.makeelement(qn("a:alpha"), {"val": _alpha_val(opacity)}))
        return sh

    def oval(self, s, x, y, w, h, fill="accent", line=None):
        sh = s.shapes.add_shape(MSO_SHAPE.OVAL, emu(x), emu(y), emu(w), emu(h))
        sh.shadow.inherit = False
        self._fill(sh, fill)
        self._line(sh, line, 1.0, False)
        return sh

    def line(self, s, x1, y1, x2, y2, color="muted", width=1.5, dash=False, arrow_end=False):
        """Straight connector from (x1, y1) to (x2, y2); `arrow_end` adds an arrowhead at (x2, y2)."""
        cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
        self._line(cn, color, width, dash)
        if arrow_end:
            cn.line._get_or_add_ln().append(parse_xml(f'<a:tailEnd {nsdecls("a")} type="triangle" w="med" len="med"/>'))
        return cn

    def arrow(self, s, x, y, w, h, color="accent", direction="right"):
        """Block arrow pointing `direction` (right, left, up, down)."""
        sh = s.shapes.add_shape(_pick("direction", direction, ARROWS), emu(x), emu(y), emu(w), emu(h))
        sh.shadow.inherit = False
        self._fill(sh, color)
        self._line(sh, None, 1.0, False)
        return sh

    # ── text ──
    def text(self, s, x, y, w, h, content, size=14, color="text", font="body", bold=False, italic=False,
             align="left", anchor="top", line_spacing=1.05):
        """Text box. `content`: str ("\\n" splits paragraphs) or a list of paragraphs; a paragraph is a
        str or a list of runs; a run is a str or {text, size?, color?, bold?, italic?, font?}."""
        sh = self._textbox(s, x, y, w, h, anchor)
        tf = sh.text_frame
        base = {"size": size, "color": color, "font": font, "bold": bold, "italic": italic}
        paras = []
        for para in [content] if isinstance(content, str) else content:
            paras.extend(para.split("\n") if isinstance(para, str) else [para])
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = _pick("align", align, ALIGN)
            p.line_spacing = line_spacing
            for run in [para] if isinstance(para, str) else para:
                if isinstance(run, str):
                    run = {"text": run}
                unknown = set(run) - RUN_KEYS
                if unknown or "text" not in run:
                    raise ValueError(f"run {run!r}: needs 'text'; valid keys: {', '.join(sorted(RUN_KEYS))}")
                self._run(p, **{**base, **run})
        return sh

    def bullets(self, s, x, y, w, h, items, size=15, marker_color="highlight", spread=True):
        """Bulleted list. "Lead: rest" bolds the lead in primary; a leading "  " makes a sub-bullet.
        `spread` distributes the free height between items (capped at 3x the font size)."""
        sh = self._textbox(s, x, y, w, h, "top")
        tf = sh.text_frame
        extra = 0.0
        if spread and len(items) > 1:
            used = 0.0
            for it in items:
                sub = it.startswith("  ")
                sz = size - (2 if sub else 0)
                used += est_lines(it.strip(), sz, w) * sz * 1.25 + (4 if sub else 10)
            extra = max(0.0, min((h * 72 - used) / (len(items) - 1), size * 3.0))
        for i, it in enumerate(items):
            sub = it.startswith("  ")
            txt = it.strip()
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt((4 if sub else 10) + extra)
            p.line_spacing = 1.05
            self._run(p, ("      " if sub else "") + "▪  ", size - (2 if sub else 0),
                      "muted" if sub else marker_color, bold=True)
            if ": " in txt and not sub:
                lead, rest = txt.split(": ", 1)
                self._run(p, lead + ": ", size, "primary", bold=True)
                self._run(p, rest, size, "text")
            else:
                self._run(p, txt, size - (1 if sub else 0), "muted" if sub else "text")
        return sh

    def chip(self, s, x, y, label, fill="accent", color="on_primary", size=9, w=None, h=0.26):
        """Pill-shaped label; `w=None` sizes it from the label length."""
        if w is None:
            w = 0.16 + 0.011 * size * len(label)
        sh = self.rect(s, x, y, w, h, fill=fill, radius=0.35)
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        self._run(p, label, size, color, bold=True)
        return sh

    # ── images and effects ──
    def image(self, s, path, x, y, w, h, fit="contain", caption=None, frame=False):
        """Picture in the box (x, y, w, h). `fit`: contain (whole image, centered), cover (cropped to
        fill), stretch. A caption takes the bottom 0.45 in of the box. GIFs stay animated."""
        path = Path(path)
        if path.suffix.lower() == ".webp":
            raise ValueError(f"{path}: .webp no soportado por python-pptx: convertir a PNG")
        if not path.is_file():
            raise FileNotFoundError(f"image not found: {path}")
        _pick("fit", fit, dict.fromkeys(FIT_MODES))
        bx, by, bw, bh = emu(x), emu(y), emu(w), emu(h - (CAPTION_H if caption else 0))
        px_w, px_h = PptxImage.from_file(str(path)).size
        if fit == "contain":
            k = min(bw / px_w, bh / px_h)
            iw, ih = int(px_w * k), int(px_h * k)
            pic = s.shapes.add_picture(str(path), bx + (bw - iw) // 2, by + (bh - ih) // 2, iw, ih)
        else:
            pic = s.shapes.add_picture(str(path), bx, by, bw, bh)
            if fit == "cover":
                img_r, box_r = px_w / px_h, bw / bh
                if img_r > box_r:
                    pic.crop_left = pic.crop_right = (1 - box_r / img_r) / 2
                else:
                    pic.crop_top = pic.crop_bottom = (1 - img_r / box_r) / 2
        left, top = pic.left / EMU_PER_IN, pic.top / EMU_PER_IN
        width, height = pic.width / EMU_PER_IN, pic.height / EMU_PER_IN
        if frame:
            self.rect(s, left, top, width, height, fill=None, line="card_border", line_w=1.25)
        if caption:
            self.text(s, x, top + height + CAPTION_GAP, w, CAPTION_H - CAPTION_GAP, caption,
                      size=11, color="muted", italic=True, align="center", line_spacing=1.0)
        return pic

    def icon(self, s, path, x, y, size=0.5):
        """Image contain-fitted in a size x size box."""
        return self.image(s, path, x, y, size, size)

    def scrim(self, s, x, y, w, h, color, opacity):
        """Semi-transparent rectangle (opacity 0-100) to dim what is behind it."""
        return self.rect(s, x, y, w, h, fill=color, opacity=opacity)

    def gradient(self, s, x, y, w, h, color, a0, a1, angle=0):
        """Rectangle whose opacity goes linearly from a0 to a1 (0-100); angle in degrees, 0 = left to right."""
        sh = self.rect(s, x, y, w, h, fill=color)
        spPr = sh._element.spPr
        solid = spPr.find(qn("a:solidFill"))
        hexv = str(self.color(color))
        stops = "".join(
            f'<a:gs pos="{pos}"><a:srgbClr val="{hexv}"><a:alpha val="{_alpha_val(a, name)}"/></a:srgbClr></a:gs>'
            for pos, a, name in ((0, a0, "a0"), (100000, a1, "a1"))
        )
        grad = parse_xml(
            f'<a:gradFill {nsdecls("a")} rotWithShape="1"><a:gsLst>{stops}</a:gsLst>'
            f'<a:lin ang="{int(angle * 60000) % 21600000}" scaled="1"/></a:gradFill>'
        )
        solid.addprevious(grad)
        spPr.remove(solid)
        return sh

    # ── tables ──
    def table(self, s, x, y, w, rows, header=True, col_widths=None, size=12):
        """Native table; the first row is the header (primary fill, on_primary bold text) when `header`."""
        if not rows or not rows[0] or any(len(r) != len(rows[0]) for r in rows):
            raise ValueError("rows must be a non-empty list of equal-length lists")
        ncols = len(rows[0])
        widths = col_widths or [w / ncols] * ncols
        if len(widths) != ncols:
            raise ValueError(f"col_widths has {len(widths)} values for {ncols} columns")
        row_h = size * 2.2 / 72
        gf = s.shapes.add_table(len(rows), ncols, emu(x), emu(y), emu(sum(widths)), emu(row_h * len(rows)))
        tbl = gf.table
        tbl.first_row = header
        tbl.horz_banding = False
        for i, cw in enumerate(widths):
            tbl.columns[i].width = emu(cw)
        for r, row in enumerate(rows):
            head = header and r == 0
            for c, value in enumerate(row):
                cell = tbl.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.color("primary" if head else "surface")
                cell.text_frame.word_wrap = True
                self._run(cell.text_frame.paragraphs[0], str(value), size,
                          "on_primary" if head else "text", bold=head)
        return gf

    # ── helpers ──
    def _fill(self, sh, color):
        if color is None:
            sh.fill.background()
        else:
            sh.fill.solid()
            sh.fill.fore_color.rgb = self.color(color)

    def _line(self, sh, color, width, dash):
        if color is None:
            sh.line.fill.background()
            return
        sh.line.color.rgb = self.color(color)
        sh.line.width = Pt(width)
        if dash:
            sh.line.dash_style = MSO_LINE.DASH

    def _textbox(self, s, x, y, w, h, anchor):
        """Wrapping text box with no insets and no autofit (the box keeps its size)."""
        sh = s.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
        tf = sh.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = _pick("anchor", anchor, ANCHOR)
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        return sh

    def _run(self, p, text, size, color="text", bold=False, italic=False, font="body"):
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = self.font(font)
        r.font.color.rgb = self.color(color)
        return r
