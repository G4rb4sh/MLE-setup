"""Components mixin for Deck: slide factories and composed content blocks.

Components draw only through the Primitives methods. Colors are brand tokens (or
colors derived from them), sizes follow ppt_lib, and the look of cards and nodes
comes from `brand.components` (card kind, radius, accent strip, node shape).
"""

from __future__ import annotations

from pathlib import Path

from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.parts.image import Image as PptxImage

from brand import NODE_SHAPES, BrandError
from deck_prims import CAPTION_H, EMU_PER_IN, est_lines

NODE_STATES = ("done", "planned")
VALIGNS = ("bias", "top")
TITLE_PLACEHOLDERS = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
KICKER_MAX_H = 1.0  # inches: tallest body placeholder taken as a kicker slot on content slides


def _gap(pt: float) -> list[dict]:
    """Empty paragraph `pt` points tall: vertical spacing inside a text box."""
    return [{"text": "", "size": pt}]


def _aspect(path) -> float:
    """Width / height of an image file (same .webp message as Primitives.image)."""
    if Path(path).suffix.lower() == ".webp":
        raise ValueError(f"{path}: .webp no soportado por python-pptx: convertir a PNG")
    w, h = PptxImage.from_file(str(path)).size
    return w / h


def _in_header(ph, title) -> bool:
    """Small placeholder (<= KICKER_MAX_H) whose bottom is not below the title's: above or beside it."""
    return ph.height <= KICKER_MAX_H * EMU_PER_IN and ph.top + ph.height <= title.top + title.height


class Components:
    # ── layout and derived colors ──
    @property
    def _m(self) -> float:
        return self.brand.grid["margin"]

    @property
    def _body_bottom(self) -> float:
        return self.H - 0.8

    def _place(self, top: float, height: float, valign: str) -> float:
        """Block top: "bias" lowers a lone block by 30% of the free space below it; "top" keeps `top`."""
        if valign not in VALIGNS:
            raise ValueError(f"valign {valign!r} is not one of {', '.join(VALIGNS)}")
        return top + max(0.0, (self._body_bottom - top - height) * 0.3) if valign == "bias" else top

    def _mix(self, a: str, b: str, t: float) -> str:
        """Hex between colors `a` (t=0) and `b` (t=1)."""
        ca, cb = self.color(a), self.color(b)
        return "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(ca, cb))

    def _lum(self, c: str) -> float:
        ch = [v / 255 for v in self.color(c)]
        r, g, b = (v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4 for v in ch)
        return 0.2126 * r + 0.7152 * g + 0.0722 * b

    def _ink(self, bg: str) -> str:
        """Heading color on `bg`: primary, or on_primary when primary would not stand out (dark styles)."""
        a, b = sorted((self._lum("primary"), self._lum(bg)))
        return "primary" if (b + 0.05) / (a + 0.05) >= 3 else "on_primary"

    # ── slide factories ──
    def _layout_slide(self, role, title, *texts, header_only=False):
        """Slide on the template layout mapped to `role` (None if unmapped).

        The title goes into the title placeholder (named "title"); each non-empty text in `texts`
        fills the next subtitle placeholder, then body placeholders top to bottom. With
        `header_only` (the caller draws below the title) a body placeholder only takes text if it
        is small and not lower than the title; texts without a slot are dropped. Placeholders
        left empty are removed so the slide shows no "Click to add" prompts.
        """
        layout = self.layout(role)
        if layout is None:
            return None
        if not any(ph.placeholder_format.type in TITLE_PLACEHOLDERS for ph in layout.placeholders):
            raise BrandError(f"deck.layouts.{role}: layout {layout.name!r} has no title placeholder")
        s = self.prs.slides.add_slide(layout)
        self.page += 1
        phs = list(s.placeholders)
        t = next(ph for ph in phs if ph.placeholder_format.type in TITLE_PLACEHOLDERS)
        t.name = "title"
        t.text_frame.text = title
        slots = sorted(
            (ph for ph in phs if ph.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE
             or ph.placeholder_format.type == PP_PLACEHOLDER.BODY and (not header_only or _in_header(ph, t))),
            key=lambda ph: (ph.placeholder_format.type != PP_PLACEHOLDER.SUBTITLE, ph.top),
        )
        for ph, text in zip(slots, [x for x in texts if x]):
            ph.text_frame.text = text
        for ph in phs:
            if not (ph.has_text_frame and ph.text_frame.text.strip()):
                ph._element.getparent().remove(ph._element)
        return s

    def title_slide(self, kicker, title, subtitle):
        s = self._layout_slide("title", title, subtitle, kicker)
        if s is not None:
            return s
        s = self.blank_slide("primary")
        w = self.W - 2.1
        self.rect(s, 0, 0, 0.28, self.H, fill="highlight")
        self.text(s, 1.1, 1.75, w, 0.4, kicker.upper(), size=15, color="highlight", font="heading", bold=True)
        self.text(s, 1.05, 2.2, w, 1.7, title, size=46, color="on_primary", font="heading", bold=True,
                  line_spacing=1.02)
        self.text(s, 1.1, 4.15, w, 0.6, subtitle, size=18, color="on_primary")
        if self.footer_text:
            self.text(s, 1.1, self.H - 0.8, w, 0.4, self.footer_text, size=12, color="on_primary",
                      font="heading", bold=True)
        return s

    def section_slide(self, number, title, image=None, caption=None, kicker="SECCIÓN"):
        """Section divider; the kicker reads "{kicker} {number}"; with `image` the big number is replaced
        by the picture."""
        label = f"{kicker} {number}".strip()
        s = self._layout_slide("section", title, label)
        if s is not None:
            if image:
                n0 = len(s.shapes)
                self.image(s, image, self.W - 4.83, 0.9, 4.2, 5.7, caption=caption)
                tree = s.shapes._spTree  # behind the layout's text: children 0-1 are the group's own props
                for el in reversed([sh._element for sh in list(s.shapes)[n0:]]):
                    tree.insert(2, el)
            return s
        s = self.blank_slide("primary")
        tw = 7.4
        if image:
            self.image(s, image, self.W - 4.83, 0.9, 4.2, 5.7, caption=caption, caption_color=self._ink("primary"))
        else:
            self.text(s, self.W - 5.13, 0.95, 5.0, 5.6, number, size=320, font="heading", bold=True,
                      color=self._mix("primary", "on_primary", 0.12), align="right", anchor="middle")
        self.text(s, self._m, 2.55, tw, 0.4, label, size=14, color="highlight",
                  font="heading", bold=True)
        self.rect(s, self._m, 3.05, 1.0, 0.09, fill="highlight")
        self.text(s, self._m, 3.3, tw, 1.6, title, size=44, color="on_primary", font="heading", bold=True)
        return s

    def content_slide(self, kicker, title):
        """Slide with header and footer (template layout: the master's); returns (slide, body top in inches)."""
        s = self._layout_slide("content", title, kicker, header_only=True)
        if s is not None:  # footer and page number come from the template master, not drawn
            return s, max(ph.top + ph.height for ph in s.placeholders) / EMU_PER_IN + 0.25
        s = self.blank_slide()
        top = self.header(s, kicker, title)
        self.footer(s)
        return s, top

    def final_slide(self, title, subtitle):
        s = self._layout_slide("final", title, subtitle)
        if s is not None:
            return s
        s = self.blank_slide("primary")
        w = self.W - 2 * self._m
        self.rect(s, 0, 0, 0.28, self.H, fill="highlight")
        self.text(s, self._m, 2.9, w, 1.4, title, size=54, color="on_primary", font="heading", bold=True,
                  anchor="middle")
        self.text(s, self._m, 4.35, w, 0.7, subtitle, size=22, color="highlight")
        if self.footer_text:
            self.text(s, self._m, self.H - 0.8, w, 0.4, self.footer_text, size=12, color="on_primary",
                      font="heading", bold=True)
        return s

    def image_slide(self, kicker, title, path, caption=None):
        """Content slide with one large centered image."""
        s, top = self.content_slide(kicker, title)
        self.image(s, path, self._m, top, self.W - 2 * self._m, self._body_bottom - top,
                   caption=caption, frame=True)
        return s

    # ── header and footer ──
    def header(self, s, kicker, title, lead=None):
        """Kicker, title (shape named "title"), optional lead and a short rule; returns the body top."""
        bg = self.brand.components["content_bg"]
        w = self.W - 2 * self._m
        self.text(s, self._m, 0.5, w, 0.3, kicker.upper(), size=12, color="accent", font="heading", bold=True)
        t = self.text(s, self._m, 0.82, w, 0.75, title, size=30, color=self._ink(bg), font="heading", bold=True)
        t.name = "title"
        y = 1.62
        if lead:
            h = est_lines(lead, 15, w) * 15 * 1.25 / 72
            self.text(s, self._m, y, w, h, lead, size=15)
            y += h + 0.12
        self.rect(s, self._m, y, 0.9, 0.07, fill="highlight")
        return y + 0.38

    def footer(self, s, page: int | None = None):
        """Brand footer (or the Deck(footer=...) override) and the page number (default: current page)."""
        y = self.H - 0.48
        if self.footer_text:
            self.text(s, self._m, y, self.W - 2 * self._m - 1.0, 0.3, self.footer_text, size=9, color="muted")
        self.text(s, self.W - self._m - 0.7, y, 0.7, 0.3, f"{self.page if page is None else page:02d}",
                  size=9, color="muted", align="right")

    # ── cards ──
    def _card_body(self, s, x, y, w, h, accent="accent"):
        """Card background in the brand's style (components.card); the shape is named "card"."""
        kind = self.brand.components["card"]
        line, line_w = {"filled": ("card_border", 1.0), "outlined": (accent, 1.5), "dark": (None, 1.0)}[kind]
        sh = self.rect(s, x, y, w, h, fill="surface", line=line, line_w=line_w,
                       radius=self.brand.components["radius"])
        sh.name = "card"
        return sh

    def _card(self, s, x, y, w, h, title, lines, accent="accent", tag=None, icon=None):
        """One card: body, accent strip (components.accent_strip), optional icon, title, lines and tag."""
        self._card_body(s, x, y, w, h, accent)
        strip = self.brand.components["accent_strip"]
        pad_l, pad_t = 0.3, 0.2
        if strip == "top":
            self.rect(s, x, y, w, 0.12, fill=accent, radius=0.5)
            pad_t += 0.12
        elif strip == "left":
            self.rect(s, x, y, 0.08, h, fill=accent, radius=0.5)
            pad_l += 0.08
        ty = y + pad_t
        if icon:
            self.icon(s, icon, x + pad_l, ty, 0.45)
            ty += 0.55
        paras = [[{"text": title, "size": 16, "bold": True, "font": "heading", "color": self._ink("surface")}],
                 _gap(6)]
        for i, ln in enumerate(lines):
            if i:
                paras.append(_gap(4))
            paras.append([{"text": "▪  ", "size": 11, "color": accent, "bold": True}, {"text": ln, "size": 12}])
        if tag:
            paras += [_gap(10), [{"text": tag, "size": 10, "color": "muted", "italic": True}]]
        self.text(s, x + pad_l, ty, w - pad_l - 0.3, y + h - 0.15 - ty, paras, anchor="middle")

    def cards_row(self, s, top, cards, height=3.2, valign="bias"):
        """Cards side by side. Items: {title, lines, accent?, tag?, icon?}. Returns the row's bottom y;
        stack rows with valign="top" (see _place)."""
        y = self._place(top, height, valign)
        n, gap = len(cards), 0.3
        cw = (self.W - 2 * self._m - gap * (n - 1)) / n
        for i, c in enumerate(cards):
            self._card(s, self._m + i * (cw + gap), y, cw, height, c["title"], c["lines"],
                      c.get("accent", "accent"), c.get("tag"), c.get("icon"))
        return y + height

    # ── content components ──
    def stat_row(self, s, top, stats):
        """Big-number panels. Items: {big, cap}."""
        n, gap = len(stats), 0.4
        cw = (self.W - 2 * self._m - gap * (n - 1)) / n
        for i, st in enumerate(stats):
            x = self._m + i * (cw + gap)
            self.rect(s, x, top, cw, 2.5, fill="primary", radius=0.05)
            self.text(s, x, top + 0.3, cw, 1.25, st["big"], size=66, color="highlight", font="heading",
                      bold=True, align="center", anchor="middle")
            self.text(s, x + 0.3, top + 1.6, cw - 0.6, 0.8, st["cap"], size=13, color="on_primary",
                      align="center")

    def _items(self, items, size, marker, color):
        paras = []
        for i, it in enumerate(items):
            if i:
                paras.append(_gap(8))
            paras.append([{"text": "▪  ", "size": size, "color": marker, "bold": True},
                          {"text": it, "size": size, "color": color}])
        return paras

    def compare(self, s, top, ltitle, litems, rtitle, ritems, valign="bias"):
        """Two panels: a neutral "before" on the left, a highlighted "after" on the right; returns the
        bottom y."""
        gap, h = 0.4, 3.7
        cw = (self.W - 2 * self._m - gap) / 2
        top = self._place(top, h, valign)
        sides = ((self._m, ltitle, litems, "surface", "card_border", "muted", "muted", "text"),
                 (self._m + cw + gap, rtitle, ritems, "primary", None, "accent", "highlight", "on_primary"))
        for x, title, items, fill, line, head, marker, color in sides:
            self.rect(s, x, top, cw, h, fill=fill, line=line, radius=0.05)
            self.rect(s, x, top, cw, 0.6, fill=head, radius=0.05)
            self.text(s, x + 0.3, top + 0.08, cw - 0.6, 0.45, title, size=15, color="on_primary",
                      font="heading", bold=True, anchor="middle")
            self.text(s, x + 0.3, top + 0.75, cw - 0.6, h - 0.95, self._items(items, 13, marker, color),
                      anchor="middle")
        return top + h

    def timeline(self, s, top, steps, box_h=2.5):
        """Numbered steps in a row. Steps: {t, a} (title, annotation)."""
        n, gap, bd = len(steps), 0.22, 0.62
        cw = (self.W - 2 * self._m - gap * (n - 1)) / n
        ink = self._ink("surface")
        for i, st in enumerate(steps):
            x = self._m + i * (cw + gap)
            self._card_body(s, x, top, cw, box_h)
            self.oval(s, x + (cw - bd) / 2, top - 0.31, bd, bd, fill="accent")
            self.text(s, x + (cw - bd) / 2, top - 0.31, bd, bd, str(i + 1), size=18, color="on_primary",
                      font="heading", bold=True, align="center", anchor="middle")
            paras = [[{"text": st["t"], "size": 13, "bold": True, "font": "heading", "color": ink}], _gap(10),
                     [{"text": st["a"], "size": 11, "bold": True, "font": "heading", "color": "accent"}]]
            self.text(s, x + 0.18, top + 0.42, cw - 0.36, box_h - 0.58, paras, align="center", anchor="middle",
                      line_spacing=1.0)

    def chain(self, s, x, y, w, h, steps, accent="accent", size=12):
        """Boxes joined by arrows: [A] → [B] → [C]. Steps: {t, a?}."""
        n, arrow = len(steps), 0.32
        bw = (w - arrow * (n - 1)) / n
        ink = self._ink("surface")
        for i, st in enumerate(steps):
            cx = x + i * (bw + arrow)
            self._card_body(s, cx, y, bw, h, accent)
            self.rect(s, cx, y, 0.08, h, fill=accent, radius=0.5)
            paras = [[{"text": st["t"], "size": size, "bold": True, "font": "heading", "color": ink}]]
            if st.get("a"):
                paras += [_gap(4), [{"text": st["a"], "size": size - 2, "color": "muted"}]]
            self.text(s, cx + 0.24, y + 0.05, bw - 0.4, h - 0.1, paras, anchor="middle", line_spacing=1.0)
            if i < n - 1:
                self.text(s, cx + bw, y, arrow, h, "→", size=17, color="muted", bold=True, align="center",
                          anchor="middle")

    def maprow(self, s, top, rows):
        """Rows of «chip → text». Rows: (chip, lead, rest); returns the y after the last row."""
        h, gap, chip_w, size = 0.66, 0.14, 2.55, 13
        ink = self._ink(self.brand.components["content_bg"])
        y = top
        for name, lead, rest in rows:
            self.rect(s, self._m, y, chip_w, h, fill="primary", radius=0.14)
            self.text(s, self._m + 0.16, y, chip_w - 0.32, h, name, size=size - 1, color="highlight",
                      font="heading", bold=True, align="center", anchor="middle", line_spacing=1.0)
            ax = self._m + chip_w
            self.text(s, ax, y, 0.34, h, "→", size=15, color="muted", bold=True, align="center", anchor="middle")
            runs = [{"text": lead, "size": size, "bold": True, "font": "heading", "color": ink}]
            if rest:
                runs.append({"text": "  " + rest, "size": size - 1, "color": "muted"})
            tx = ax + 0.34
            self.text(s, tx, y, self.W - self._m - tx, h, [runs], anchor="middle")
            y += h + gap
        return y

    def numbered_grid(self, s, top, items, height=None):
        """Two columns of numbered items. Items: (number, label)."""
        gap, bd = 0.6, 0.5
        cw = (self.W - 2 * self._m - gap) / 2
        rows = (len(items) + 1) // 2
        rh = min(max((height or (self._body_bottom - top)) / rows, 0.82), 1.3)
        for idx, (num, label) in enumerate(items):
            x = self._m + (idx // rows) * (cw + gap)
            y = top + (idx % rows) * rh + (rh - bd) / 2
            self.oval(s, x, y, bd, bd, fill="primary")
            self.text(s, x, y, bd, bd, num, size=14, color="highlight", font="heading", bold=True,
                      align="center", anchor="middle")
            self.text(s, x + bd + 0.22, y, cw - bd - 0.22, bd, label, size=14, anchor="middle", line_spacing=1.0)

    def code_block(self, s, top, lines, height=3.6, note=None, valign="bias"):
        """Monospaced block on a primary panel; lines starting with "# " are highlighted. Returns the
        bottom y (below the note, if any)."""
        w = self.W - 2 * self._m
        y = self._place(top, height, valign)
        self.rect(s, self._m, y, w, height, fill="primary", radius=0.03)
        paras = [[{"text": ln, "color": "highlight" if ln.startswith("# ") else "on_primary"}] for ln in lines]
        self.text(s, self._m + 0.45, y + 0.32, w - 0.9, height - 0.6, paras, size=13, font="mono",
                  line_spacing=1.25)
        if note:
            self.text(s, self._m, y + height + 0.12, w, 0.35, note, size=11, color="muted", italic=True)
            return y + height + 0.47
        return y + height

    # ── images ──
    def image_row(self, s, top, items, height=None):
        """Images in a row scaled to a common height so captions align. Items: {path, caption?}."""
        total, gap = self.W - 2 * self._m, 0.35
        cap_h = CAPTION_H if any(it.get("caption") for it in items) else 0.0
        box_h = (height or (self._body_bottom - top)) - cap_h
        ars = [_aspect(it["path"]) for it in items]
        h = min(box_h, (total - gap * (len(items) - 1)) / sum(ars))
        widths = [ar * h for ar in ars]
        x = self._m + (total - sum(widths) - gap * (len(items) - 1)) / 2
        for it, w in zip(items, widths):
            self.image(s, it["path"], x, top, w, h + cap_h, caption=it.get("caption"), frame=True)
            x += w + gap

    def split_image(self, s, top, path, items, side="right", ratio=0.4, caption=None):
        """Bullets on one side, an image on the other; `ratio` is the image's share of the width."""
        if side not in ("left", "right"):
            raise ValueError(f"side {side!r} is not one of left, right")
        gap = 0.5
        total = self.W - 2 * self._m - gap
        iw = total * ratio
        tw = total - iw
        h = self._body_bottom - top
        tx, ix = (self._m, self._m + tw + gap) if side == "right" else (self._m + iw + gap, self._m)
        self.bullets(s, tx, top, tw, h, items, lead_color=self._ink(self.brand.components["content_bg"]))
        self.image(s, path, ix, top, iw, h, caption=caption, frame=True)

    # ── diagram pieces ──
    def _node_look(self, state, accent):
        """(fill, line, dash) for a done or planned node."""
        if state not in NODE_STATES:
            raise ValueError(f"state {state!r} is not one of {', '.join(NODE_STATES)}")
        if state == "done":
            return "surface", accent, False
        return self._mix("surface", self.brand.components["content_bg"], 0.5), "planned", True

    def node(self, s, x, y, w, h, label, sub=None, accent="accent", state="done", shape=None):
        """Diagram box. `planned` draws a dashed border and a dimmer fill; `shape` is rounded or rect
        (default components.node_shape). Returns the box shape."""
        shape = shape or self.brand.components["node_shape"]
        if shape not in NODE_SHAPES:
            raise ValueError(f"shape {shape!r} is not one of {', '.join(NODE_SHAPES)}")
        fill, line, dash = self._node_look(state, accent)
        box = self.rect(s, x, y, w, h, fill=fill, line=line, line_w=1.4, dash=dash,
                        radius=self.brand.components["radius"] if shape == "rounded" else None)
        color = self._ink("surface") if state == "done" else "text"
        paras = [[{"text": label, "size": 11.5, "bold": True, "font": "heading", "color": color}]]
        if sub:
            paras.append([{"text": sub, "size": 8.6, "color": "muted"}])
        self.text(s, x + 0.08, y, w - 0.16, h, paras, align="center", anchor="middle", line_spacing=1.08)
        return box

    def labeled_arrow(self, s, x, y, w, label=None, color="accent"):
        """Right arrow from x to x+w at height y, with an optional label above it (as wide as the arrow)."""
        self.arrow(s, x, y + 0.02, w, 0.2, color=color)
        if label:
            self.text(s, x, y - 0.24, w, 0.2, label, size=8.2, color="muted", font="heading", bold=True,
                      align="center")

    def container(self, s, x, y, w, h, label):
        """Grouping box with a small label at the top-left, behind the nodes drawn after it."""
        bg = self.brand.components["content_bg"]
        self.rect(s, x, y, w, h, fill=self._mix(bg, "surface", 0.5), line="card_border",
                  radius=self.brand.components["radius"] / 2)
        self.text(s, x + 0.3, y + 0.18, w - 0.6, 0.22, label, size=9.5, color="muted", font="heading", bold=True)

    def legend(self, s, x, y, items=None):
        """Node-state key. Items: (label, "done" | "planned")."""
        for label, state in items or [("Hecho", "done"), ("Planeado", "planned")]:
            fill, line, dash = self._node_look(state, "accent")
            self.rect(s, x, y + 0.01, 0.2, 0.17, fill=fill, line=line, line_w=1.4, dash=dash, radius=0.3)
            tw = len(label) * 10 * 0.6 / 72 + 0.1
            self.text(s, x + 0.32, y - 0.03, tw, 0.24, label, size=10, bold=True)
            x += 0.32 + tw + 0.4
