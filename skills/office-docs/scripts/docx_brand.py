"""Brand look for pandoc-generated .docx files: reference styles, cover, tables, header/footer, TOC.

Every color and font comes from the brand. Color roles and their fallbacks:
heading = docx.heading_color | colors.accent, cover band = docx.cover_band | colors.primary,
code = docx.code_bg | F2F2F2, notes = docx.note_bg | colors.background,
footer rule = docx.rule | colors.card_border; text = colors.text, muted = colors.muted.
Styles live in the reference document (real Word styles); only the cover, tables and
header/footer get direct formatting.
"""

from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from brand import Brand
from render import office_app, office_script, permission_hint, run_office_script, temp_copy_name

WORD = "Microsoft Word"
TOC_TIMEOUT_S = 300
EMU_PER_IN = 914400

LABELS = {
    "es": {"for": "Preparado para", "by": "Preparado por", "version": "Versión",
           "page": "Página", "toc": "Tabla de contenido"},
    "en": {"for": "Prepared for", "by": "Prepared by", "version": "Version",
           "page": "Page", "toc": "Table of contents"},
}


def labels(lang: str) -> dict[str, str]:
    """Fixed document words in the document's language (Spanish unless `lang` is English)."""
    return LABELS["en" if lang.lower().startswith("en") else "es"]


def palette(brand: Brand) -> dict[str, str]:
    """Hex colors for each document role (see the module docstring for the fallbacks)."""
    c, d = brand.colors, brand.docx
    return {
        "heading": d.get("heading_color", c["accent"]),
        "band": d.get("cover_band", c["primary"]),
        "code_bg": d.get("code_bg", "F2F2F2"),
        "note_bg": d.get("note_bg", c["background"]),
        "rule": d.get("rule", c["card_border"]),
        "text": c["text"],
        "muted": c["muted"],
        "accent": c["accent"],
        "on_primary": c["on_primary"],
    }


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


# --------------------------------------------------------------------------- #
# Low-level OOXML helpers

def _shade(el, fill: str) -> None:
    for old in el.findall(qn("w:shd")):
        el.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el.append(shd)


def _paragraph_border(ppr, sides: dict[str, tuple[str, int, int]]) -> None:
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        ppr.append(bdr)
    for side, (color, size, space) in sides.items():
        for old in bdr.findall(qn(f"w:{side}")):
            bdr.remove(old)
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(size))
        b.set(qn("w:space"), str(space))
        b.set(qn("w:color"), color)
        bdr.append(b)


def _run_font(rpr, name: str) -> None:
    """Set all four font slots and drop theme fonts, so the named font wins everywhere."""
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        rfonts.attrib.pop(qn(attr), None)


def _style_font(style, name: str, size: float, color: str | None = None,
                bold: bool | None = None) -> None:
    f = style.font
    f.name = name
    f.size = Pt(size)
    if color is not None:
        f.color.rgb = _rgb(color)
    if bold is not None:
        f.bold = bold
    _run_font(style.element.get_or_add_rPr(), name)


def _heading(style, font: str, color: str, size: float, before: float, after: float) -> None:
    _style_font(style, font, size, color, False)
    style.font.italic = False
    pf = style.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.0
    pf.keep_with_next = True


def _field(paragraph, instruction: str) -> None:
    for kind in ("begin", None, "separate", "end"):
        run = paragraph.add_run()
        if kind is None:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = instruction
            run._r.append(instr)
        else:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._r.append(fld)


def _page_break():
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    p.append(r)
    return p


# --------------------------------------------------------------------------- #
# Reference document (styles)

def make_reference_docx(brand: Brand, dest: Path, page_break_on_h1: bool) -> Path:
    """Write the pandoc reference .docx for `brand` to `dest`.

    With `[docx] reference` the brand's file is used as is (only Heading 1's page break is
    set); otherwise pandoc's default reference gets the brand's fonts and colors.
    """
    dest = Path(dest)
    if brand.docx.get("reference"):
        shutil.copyfile(brand.docx["reference"], dest)
        doc = Document(str(dest))
        doc.styles["Heading 1"].paragraph_format.page_break_before = page_break_on_h1
        doc.save(str(dest))
        return dest

    subprocess.run(["pandoc", "-o", str(dest), "--print-default-data-file", "reference.docx"],
                   check=True, capture_output=True)
    doc = Document(str(dest))
    s = doc.styles
    col = palette(brand)
    body, heading, mono = brand.fonts["body"], brand.fonts["heading"], brand.fonts["mono"]

    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in s:
            _style_font(s[name], body, 10.5, col["text"])
    for name in ("Body Text", "First Paragraph"):
        pf = s[name].paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(6), Pt(6), 1.15
    pf = s["Compact"].paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(2), Pt(2), 1.15

    for name, size, before, after in (("Heading 1", 18, 24, 12), ("Heading 2", 15, 18, 8),
                                      ("Heading 3", 12.5, 14, 6), ("Heading 4", 11, 10, 4),
                                      ("Title", 22, 0, 6), ("Subtitle", 13, 0, 36)):
        _heading(s[name], heading, col["heading"], size, before, after)
    s["Heading 1"].paragraph_format.page_break_before = page_break_on_h1
    for name in ("Title", "Subtitle"):  # pandoc centers them and spaces the subtitle's letters
        s[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for spacing in s[name].element.get_or_add_rPr().findall(qn("w:spacing")):
            spacing.getparent().remove(spacing)
    if "TOC Heading" in s:
        _heading(s["TOC Heading"], heading, col["heading"], 18, 0, 18)
    for level, size, indent in ((1, 11.5, 0), (2, 10, 0.15), (3, 10, 0.3)):
        name = f"toc {level}"
        style = s[name] if name in s else s.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        _style_font(style, body, size, col["text"])
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.left_indent = Inches(indent)

    # pandoc's default reference has no "Source Code"; pandoc keeps the one given here.
    code = s["Source Code"] if "Source Code" in s else s.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
    _style_font(code, mono, 8.5)
    pf = code.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(4), Pt(8), 1.0
    ppr = code.element.get_or_add_pPr()
    _paragraph_border(ppr, {side: (col["code_bg"], 4, 6) for side in ("top", "left", "bottom", "right")})
    _shade(ppr, col["code_bg"])
    if "Verbatim Char" in s:
        _style_font(s["Verbatim Char"], mono, 9.5, col["text"])
        _shade(s["Verbatim Char"].element.get_or_add_rPr(), col["code_bg"])

    # Notes (block quotes): accent bar on the left over a light background.
    note = s["Block Text"]
    _style_font(note, body, 10, col["text"])
    note.font.italic = False
    ppr = note.element.get_or_add_pPr()
    _paragraph_border(ppr, {"left": (col["accent"], 24, 10)})
    _shade(ppr, col["note_bg"])
    note.paragraph_format.left_indent = Inches(0.15)
    note.paragraph_format.right_indent = Inches(0.1)

    if "Hyperlink" in s:
        s["Hyperlink"].font.color.rgb = _rgb(col["accent"])

    # Letter, one-inch margins, header and footer close to the edge.
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance, sec.footer_distance = Inches(0.49), Inches(0.3)
    doc.save(str(dest))
    return dest


# --------------------------------------------------------------------------- #
# Tables

def _table_borders(table, color: str) -> None:
    """Horizontal lines only, in the heading color; full width; roomy cells."""
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        if side in ("left", "right", "insideV"):
            b.set(qn("w:val"), "nil")
        else:
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), color)
        borders.append(b)
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", 60), ("bottom", 60), ("left", 100), ("right", 100)):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), str(value))
        m.set(qn("w:type"), "dxa")
        margins.append(m)
    for tag, new in (("w:tblBorders", borders), ("w:tblCellMar", margins)):
        for old in tblpr.findall(qn(tag)):
            tblpr.remove(old)
        tblpr.append(new)
    width = tblpr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        tblpr.append(width)
    width.set(qn("w:type"), "pct")
    width.set(qn("w:w"), "5000")


def style_tables(doc, brand: Brand) -> None:
    """Header row filled with the heading color (repeated on each page), horizontal rules only."""
    col = palette(brand)
    font = brand.fonts["heading"]
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(table, col["heading"])
        for i, row in enumerate(table.rows):
            if i == 0:
                trpr = row._tr.get_or_add_trPr()
                for old in trpr.findall(qn("w:tblHeader")):
                    trpr.remove(old)
                repeat = OxmlElement("w:tblHeader")
                repeat.set(qn("w:val"), "true")
                trpr.append(repeat)
            for cell in row.cells:
                if i == 0:
                    _shade(cell._tc.get_or_add_tcPr(), col["heading"])
                for p in cell.paragraphs:
                    pf = p.paragraph_format
                    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(0), 1.0
                    for r in p.runs:
                        r.font.size = Pt(9.5)
                        if i == 0:
                            r.font.name = font
                            r.font.bold = True
                            _run_font(r._r.get_or_add_rPr(), font)
                            r.font.color.rgb = _rgb(col["on_primary"])


# --------------------------------------------------------------------------- #
# Header and footer

def _muted(paragraph, text: str, brand: Brand, size: float = 8) -> None:
    r = paragraph.add_run(text)
    r.font.name, r.font.size, r.font.color.rgb = brand.fonts["body"], Pt(size), _rgb(palette(brand)["muted"])


def add_header_footer(doc, meta, brand: Brand) -> None:
    """Header: light logo (if any) and the title. Footer: a rule, client and version, page number.

    The first page (the cover) has neither.
    """
    words = labels(meta.lang)
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    width = sec.page_width - sec.left_margin - sec.right_margin

    head = sec.header.paragraphs[0]
    head.text = ""
    head.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    if "light" in brand.logos:
        head.add_run().add_picture(str(brand.logos["light"]), width=Inches(1.0))
    _muted(head, f"\t{meta.title}", brand)

    foot = sec.footer.paragraphs[0]
    foot.text = ""
    foot.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    _paragraph_border(foot._p.get_or_add_pPr(), {"top": (palette(brand)["rule"], 4, 6)})
    left = [f"{words['for']} {meta.client}" if meta.client else "", f"{words['version']} {meta.version}"]
    _muted(foot, f"{'  ·  '.join(x for x in left if x)}\t{words['page']} ", brand)
    _field(foot, "PAGE")
    for r in foot.runs:
        r.font.name, r.font.size, r.font.color.rgb = brand.fonts["body"], Pt(8), _rgb(palette(brand)["muted"])


# --------------------------------------------------------------------------- #
# Cover (DrawingML shapes anchored to the page)

_NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
       'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
       'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
       'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
       'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"')


def _anchor(id_: int, name: str, x: int, y: int, cx: int, cy: int, layer: int, graphic: str) -> str:
    # The band (layer 0) goes behind the text; the rest on top in layer order.
    behind = "1" if layer == 0 else "0"
    return f"""
<wp:anchor {_NS} distT="0" distB="0" distL="0" distR="0" simplePos="0"
    relativeHeight="{251658240 + layer * 1024}" behindDoc="{behind}" locked="1"
    layoutInCell="1" allowOverlap="1">
  <wp:simplePos x="0" y="0"/>
  <wp:positionH relativeFrom="page"><wp:posOffset>{x}</wp:posOffset></wp:positionH>
  <wp:positionV relativeFrom="page"><wp:posOffset>{y}</wp:posOffset></wp:positionV>
  <wp:extent cx="{cx}" cy="{cy}"/>
  <wp:effectExtent l="0" t="0" r="0" b="0"/>
  <wp:wrapNone/>
  <wp:docPr id="{id_}" name="{_escape(name)}"/>
  <wp:cNvGraphicFramePr/>
  <a:graphic>{graphic}</a:graphic>
</wp:anchor>"""


def _rectangle(cx: int, cy: int, color: str, content: str = "") -> str:
    text = f"<wps:txbx><w:txbxContent>{content}</w:txbxContent></wps:txbx>" if content else ""
    return f"""
<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
  <wps:wsp>
    <wps:cNvSpPr/>
    <wps:spPr>
      <a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>
      <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
      <a:solidFill><a:srgbClr val="{color}"/></a:solidFill>
      <a:ln><a:noFill/></a:ln>
    </wps:spPr>
    {text}
    <wps:bodyPr rot="0" vert="horz" wrap="square" lIns="457200" tIns="182880"
        rIns="228600" bIns="91440" anchor="ctr"><a:noAutofit/></wps:bodyPr>
  </wps:wsp>
</a:graphicData>"""


def _picture(rid: str, cx: int, cy: int) -> str:
    return f"""
<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">
  <pic:pic>
    <pic:nvPicPr><pic:cNvPr id="0" name="Logo"/><pic:cNvPicPr/></pic:nvPicPr>
    <pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>
    <pic:spPr>
      <a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>
      <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
    </pic:spPr>
  </pic:pic>
</a:graphicData>"""


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def _box_paragraph(text: str, size: float, color: str, font: str) -> str:
    font = _escape(font)
    return (f'<w:p><w:pPr><w:spacing w:before="0" w:after="0" w:line="240" '
            f'w:lineRule="auto"/></w:pPr><w:r><w:rPr><w:rFonts w:ascii="{font}" '
            f'w:hAnsi="{font}" w:cs="{font}"/><w:color w:val="{color}"/>'
            f'<w:sz w:val="{int(size * 2)}"/></w:rPr><w:t xml:space="preserve">'
            f'{_escape(text)}</w:t></w:r></w:p>')


def _inches(value: float) -> int:
    return int(value * EMU_PER_IN)


def _cover_elements(doc, meta, brand: Brand) -> list:
    col = palette(brand)
    words = labels(meta.lang)
    page_w = int(doc.sections[0].page_width)
    band_h = _inches(4.2)
    shapes = [_anchor(9001, "Band", 0, 0, page_w, band_h, 0, _rectangle(page_w, band_h, col["band"]))]
    if meta.project:
        box_w, box_h = _inches(4.9), _inches(2.0)
        shapes.append(_anchor(9002, "Project", 0, _inches(2.2), box_w, box_h, 1, _rectangle(
            box_w, box_h, col["accent"],
            _box_paragraph(meta.project, 20, col["on_primary"], brand.fonts["heading"]))))
    if "dark" in brand.logos:  # the band is dark: use the logo made for dark backgrounds
        rid, image = doc.part.get_or_add_image(str(brand.logos["dark"]))
        cx = _inches(1.5)
        cy = int(cx * image.px_height / image.px_width)
        shapes.append(_anchor(9003, "Logo", _inches(0.45), _inches(0.35), cx, cy, 2, _picture(rid, cx, cy)))

    anchor = doc.add_paragraph()
    for shape in shapes:
        drawing = OxmlElement("w:drawing")
        drawing.append(parse_xml(shape))
        anchor.add_run()._r.append(drawing)
    elements = [anchor._p]

    def line(text: str, size: float | None = None, color: str | None = None, before: float = 0,
             after: float = 0, style: str | None = None) -> None:
        p = doc.add_paragraph(style=style)
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.left_indent = Inches(-0.5)  # in line with the text of the project box (0.5" from the edge)
        pf.page_break_before = False
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        if size is not None:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = _rgb(color)
        elements.append(p._p)

    # The band takes 4.2" from the top edge; the top margin already leaves 1".
    line(meta.title, before=250, after=6, style="Title")
    if meta.subtitle:
        line(meta.subtitle, after=36, style="Subtitle")
    for label, value in ((words["for"], meta.client),
                         (words["by"], meta.author or brand.docx.get("company", ""))):
        if value:
            line(label, 9, col["muted"])
            line(value, 11, col["text"], after=10)
    line(f"{meta.date}  ·  {words['version']} {meta.version}", 10, col["muted"])
    return elements


def _is_h1(el) -> bool:
    style = el.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
    return style is not None and style.get(qn("w:val")) == "Heading1"


def _break_before(el, h1_breaks: bool) -> None:
    """A page break before `el`, unless `el` is a Heading 1 that already starts a page."""
    if el is None or el.tag == qn("w:sectPr") or (h1_breaks and _is_h1(el)):
        return
    el.addprevious(_page_break())


def add_cover(doc, meta, brand: Brand) -> None:
    """Cover page (band, optional project box and logo, title and data) before the TOC.

    The cover and the TOC each get their own page.
    """
    body = doc.element.body
    toc = body.find(qn("w:sdt"))
    first = toc if toc is not None else body[0]
    after_toc = toc.getnext() if toc is not None else None
    for el in _cover_elements(doc, meta, brand):
        first.addprevious(el)
    _break_before(first, meta.page_break_on_h1)
    _break_before(after_toc, meta.page_break_on_h1)


# --------------------------------------------------------------------------- #
# Fields and TOC

_SETTINGS_AFTER_UPDATE_FIELDS = (
    "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
    "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
    "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown",
    "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol",
    "listSeparator",
)


def update_fields_on_open(docx: Path) -> None:
    """Ask Word to update fields (the TOC) when the file is opened; Word asks the user first."""
    doc = Document(str(docx))
    settings = doc.settings.element
    for old in settings.findall(qn("w:updateFields")):
        settings.remove(old)
    flag = OxmlElement("w:updateFields")
    flag.set(qn("w:val"), "true")
    after = {qn(f"w:{n}") for n in _SETTINGS_AFTER_UPDATE_FIELDS}
    for child in settings:
        if child.tag in after:
            child.addprevious(flag)
            break
    else:
        settings.append(flag)
    doc.save(str(docx))


_TOC_ACTION = """update (table of contents 1 of $doc)
      save $doc"""


def _toc_script(copy: Path) -> str:
    return office_script(WORD, copy, _TOC_ACTION, timeout=TOC_TIMEOUT_S)


def fill_toc_with_word(docx: Path) -> bool:
    """Fill the TOC with Microsoft Word (macOS). True on success; False (never raises) otherwise.

    Word only ever sees a temporary copy with a unique `office-docs-<hex>.docx` name, addressed
    by that name in every Apple event; the result replaces `docx` only on success. Documents
    the user has open are never touched.
    """
    docx = Path(docx)
    if office_app(".docx") != WORD:
        return False
    try:
        with tempfile.TemporaryDirectory(prefix="office-docs-toc-") as td:
            copy = Path(td) / temp_copy_name(".docx")
            shutil.copyfile(docx, copy)
            r = run_office_script(WORD, copy, _toc_script(copy), timeout=TOC_TIMEOUT_S)
            if r.returncode != 0:
                err = r.stderr.strip()
                print(f"warning: Word could not fill the table of contents ({err}); it will be "
                      f"updated when the document is opened.", file=sys.stderr)
                hint = permission_hint(WORD, err)
                if hint:
                    print(hint, file=sys.stderr)
                return False
            shutil.copyfile(copy, docx)
    except (OSError, subprocess.SubprocessError) as e:
        print(f"warning: Word could not fill the table of contents ({e}); it will be updated "
              f"when the document is opened.", file=sys.stderr)
        return False
    return True


# --------------------------------------------------------------------------- #
# OOXML child order

# Child order the OOXML schema requires (only the parents touched here).
_ORDEN = {
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
            "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
            "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
            "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
            "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
            "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
            "sectPr", "pPrChange"],
    "rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike",
            "dstrike", "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid",
            "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs",
            "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs",
            "em", "lang", "eastAsianLayout", "specVanish", "oMath"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
             "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
              "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
              "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"],
    "trPr": ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit",
             "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden"],
}


def _order_children(root) -> None:
    for parent, order in _ORDEN.items():
        position = {qn(f"w:{n}"): i for i, n in enumerate(order)}
        for el in root.iter(qn(f"w:{parent}")):
            children = sorted(el, key=lambda c: position.get(c.tag, len(order)))
            for c in children:
                el.remove(c)
                el.append(c)


def fix_property_order(docx: Path) -> None:
    """Reorder paragraph/run/table property children as the schema requires (Word is strict)."""
    doc = Document(str(docx))
    for part in doc.part.package.iter_parts():
        if isinstance(part, XmlPart):
            _order_children(part.element)
    doc.save(str(docx))

