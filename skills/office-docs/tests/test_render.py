import os
import re
import subprocess
import sys
from pathlib import Path

import pymupdf
import pytest

import render
from conftest import SCRIPTS
from render import (RenderUnavailable, applescript_quote, export_pdf, office_sandbox_dirs, pdf_to_png,
                    run_build)
from test_components import _deck, _png, build_smoke


def _pdf(path, pages):
    doc = pymupdf.open()
    for i in range(pages):
        doc.new_page(width=400, height=300).insert_text((50, 100), f"page {i + 1}")
    doc.save(str(path))
    doc.close()
    return path


def test_pdf_to_png(tmp_path):
    out = tmp_path / "png"
    out.mkdir()
    (out / "07.png").write_bytes(b"stale")
    pngs = pdf_to_png(_pdf(tmp_path / "x.pdf", 2), out, dpi=72)
    assert [p.name for p in pngs] == ["01.png", "02.png"]
    assert sorted(p.name for p in out.iterdir()) == ["01.png", "02.png"]
    pix = pymupdf.Pixmap(str(pngs[0]))
    assert (pix.width, pix.height) == (400, 300)


def test_applescript_quote():
    q = applescript_quote(Path('/a/Overview Técnico — Q3 "x".pptx'))
    assert q.startswith('"') and q.endswith('"')
    assert '\\"x\\"' in q
    assert "Técnico — Q3" in q
    assert applescript_quote(Path("/a/b\\c")) == '"/a/b\\\\c"'


def test_run_build_sets_pythonpath(tmp_path, monkeypatch):
    monkeypatch.setenv("PYTHONPATH", "/somewhere/else")
    out = tmp_path / "env.txt"
    script = tmp_path / "build.py"
    script.write_text(
        "import os, pathlib\n"
        "pathlib.Path('env.txt').write_text(os.environ['PYTHONPATH'])\n"
        "import deck\n",
        encoding="utf-8",
    )
    run_build(script)
    pp = out.read_text().split(os.pathsep)
    assert pp[0] == str(SCRIPTS)
    assert "/somewhere/else" in pp

    bad = tmp_path / "bad.py"
    bad.write_text("raise SystemExit(3)\n", encoding="utf-8")
    with pytest.raises(subprocess.CalledProcessError):
        run_build(bad)


def test_unavailable(monkeypatch, tmp_path):
    monkeypatch.setattr(render, "office_app", lambda suffix: None)
    monkeypatch.setattr(render.shutil, "which", lambda name: None)
    monkeypatch.setattr(render.sys, "platform", "linux")
    src = tmp_path / "deck.pptx"
    src.write_bytes(b"")
    with pytest.raises(RenderUnavailable) as e:
        export_pdf(src, tmp_path / "deck.pdf")
    msg = str(e.value)
    assert "PowerPoint" in msg and "soffice" in msg


def test_office_app_detection(monkeypatch):
    monkeypatch.setattr(render.sys, "platform", "linux")
    assert render.office_app(".pptx") is None
    monkeypatch.setattr(render.sys, "platform", "darwin")
    monkeypatch.setattr(render.Path, "exists", lambda self: True)
    assert render.office_app(".pptx") == "Microsoft PowerPoint"
    assert render.office_app(".docx") == "Microsoft Word"
    assert render.office_app(".txt") is None


def test_render_defaults(monkeypatch, tmp_path):
    def fake_export(src, pdf):
        return _pdf(pdf, 3)

    monkeypatch.setattr(render, "export_pdf", fake_export)
    src = tmp_path / "deck.pptx"
    src.write_bytes(b"")
    res = render.render(src, dpi=50)
    assert res.pdf == tmp_path / "deck.pdf"
    assert [p.name for p in res.pngs] == ["01.png", "02.png", "03.png"]
    assert all(p.parent == tmp_path / "deck-render" for p in res.pngs)

    only = render.render(src, pdf_only=True)
    assert only.pngs == []


@pytest.mark.skipif(os.environ.get("OFFICE_DOCS_RENDER") != "1", reason="set OFFICE_DOCS_RENDER=1")
def test_real_export(make_brand, tmp_path):
    d = _deck(make_brand, "modern")
    build_smoke(d, _png(tmp_path / "img.png", 300, 200))
    folder = tmp_path / "Carpeta con espacio — ñ"
    folder.mkdir()
    path = folder / "deck.pptx"
    d.save(path)
    res = render.render(path, dpi=40)
    assert res.pdf.exists()
    assert len(res.pngs) >= 1


@pytest.mark.skipif(os.environ.get("OFFICE_DOCS_RENDER") != "1", reason="set OFFICE_DOCS_RENDER=1")
def test_real_export_docx(tmp_path):
    from docx import Document

    folder = tmp_path / 'Carpeta "doc" — ñ'
    folder.mkdir()
    path = folder / "doc.docx"
    doc = Document()
    doc.add_heading("Hola", 1)
    doc.add_paragraph("Un párrafo.")
    doc.save(str(path))
    res = render.render(path, dpi=40)
    assert len(res.pngs) >= 1


def test_not_authorized(monkeypatch, tmp_path):
    def fake_run(cmd, input, **kw):
        return subprocess.CompletedProcess(cmd, 1, "", "execution error: Not authorized to send "
                                           "Apple events to Microsoft PowerPoint. (-1743)")

    monkeypatch.setattr(render, "office_app", lambda suffix: "Microsoft PowerPoint")
    monkeypatch.setattr(render.subprocess, "run", fake_run)
    src = tmp_path / "deck.pptx"
    src.write_bytes(b"")
    with pytest.raises(RenderUnavailable, match="Privacy & Security → Automation"):
        export_pdf(src, tmp_path / "deck.pdf")


@pytest.mark.parametrize("app,noun", [("Microsoft PowerPoint", "presentation"),
                                      ("Microsoft Word", "document")])
def test_script_targets_copy_by_name(app, noun):
    copy = Path('/var/x/Carpeta "q" — ñ/office-docs-0a1b2c3d.pptx')
    script = render._office_script(app, copy, copy.with_suffix(".pdf"))
    assert f"set nm to {applescript_quote(Path(copy.name))}" in script
    acts = [ln.strip() for ln in script.splitlines()
            if re.match(r"\s*(close|save)\b", ln)]
    assert len(acts) == 3  # save, close, best-effort close on the late-open path
    assert all(f"{noun} nm" in ln for ln in acts)
    assert not re.search(rf"{noun} (idx|i|\d+)\b", script)
    assert "may still appear" in script


def test_main_timeout_exits_2(monkeypatch, tmp_path, capsys):
    def slow(*a, **kw):
        raise subprocess.TimeoutExpired(["osascript", "-"], 630)

    monkeypatch.setattr(render, "render", slow)
    assert render.main([str(tmp_path / "deck.pptx")]) == 2
    assert "render failed" in capsys.readouterr().err


PP_CONTAINER = Path("Library/Containers/com.microsoft.Powerpoint/Data/tmp")


def _fake_export(seen):
    """A run_office_script stand-in: records the copy and the script, writes the PDF it names."""
    def run(app, copy, script, timeout=render.EXPORT_TIMEOUT_S):
        assert copy.exists()
        seen.append((copy, script))
        Path(re.search(r'set pdfPath to "(.+)"', script).group(1)).write_bytes(b"%PDF")
        return subprocess.CompletedProcess(["osascript"], 0, "", "")
    return run


def test_copy_lives_in_app_container(monkeypatch, tmp_path):
    home = tmp_path / "home"
    (home / PP_CONTAINER).mkdir(parents=True)
    monkeypatch.setattr(render, "office_sandbox_dirs", office_sandbox_dirs)
    monkeypatch.setattr(render.Path, "home", classmethod(lambda cls: home))
    monkeypatch.setattr(render, "office_app", lambda suffix: "Microsoft PowerPoint")
    seen = []
    monkeypatch.setattr(render, "run_office_script", _fake_export(seen))
    src = tmp_path / "deck.pptx"
    src.write_bytes(b"")
    assert export_pdf(src, tmp_path / "out" / "deck.pdf").read_bytes() == b"%PDF"
    [(copy, script)] = seen
    work = copy.parent
    assert work.parent == home / PP_CONTAINER and work.name.startswith("office-docs-render-")
    assert f"set pdfPath to {applescript_quote(work / copy.with_suffix('.pdf').name)}" in script
    assert not work.exists() and list((home / PP_CONTAINER).iterdir()) == []


@pytest.mark.parametrize("state", ["missing", "denied"])
def test_copy_falls_back_to_temp(monkeypatch, tmp_path, state):
    home = tmp_path / "home"
    container = home / PP_CONTAINER
    if state == "denied":
        container.mkdir(parents=True)
        container.chmod(0o500)
    monkeypatch.setattr(render, "office_sandbox_dirs", office_sandbox_dirs)
    monkeypatch.setattr(render.Path, "home", classmethod(lambda cls: home))
    monkeypatch.setattr(render, "office_app", lambda suffix: "Microsoft PowerPoint")
    seen = []
    monkeypatch.setattr(render, "run_office_script", _fake_export(seen))
    src = tmp_path / "deck.pptx"
    src.write_bytes(b"")
    try:
        export_pdf(src, tmp_path / "deck.pdf")
    finally:
        if state == "denied":
            container.chmod(0o700)
    work = seen[0][0].parent
    assert home not in work.parents and work.name.startswith("office-docs-render-")
    assert not work.exists()


def test_failed_export_removes_workdir(monkeypatch, tmp_path):
    home = tmp_path / "home"
    (home / PP_CONTAINER).mkdir(parents=True)
    monkeypatch.setattr(render, "office_sandbox_dirs", office_sandbox_dirs)
    monkeypatch.setattr(render.Path, "home", classmethod(lambda cls: home))
    monkeypatch.setattr(render, "office_app", lambda suffix: "Microsoft PowerPoint")
    monkeypatch.setattr(render.subprocess, "run",
                        lambda cmd, **kw: subprocess.CompletedProcess(cmd, 1, "", "boom (-9074)"))
    src = tmp_path / "deck.pptx"
    src.write_bytes(b"")
    with pytest.raises(RuntimeError, match="press Esc in Microsoft PowerPoint"):
        export_pdf(src, tmp_path / "deck.pdf")
    assert list((home / PP_CONTAINER).iterdir()) == []


def test_close_error_does_not_mask_action_error():
    copy = Path("/var/x/office-docs-0a1b2c3d.pptx")
    script = render._office_script("Microsoft PowerPoint", copy, copy.with_suffix(".pdf"))
    after = script.split("set failure to {m, n}\n    end try\n", 1)[1]
    assert re.match(r"\s*try\n\s*close presentation nm saving no\n\s*on error m number n\n"
                    r"\s*if failure is missing value then set failure to \{m, n\}\n\s*end try\n"
                    r"\s*if failure is not missing value then error", after), after
