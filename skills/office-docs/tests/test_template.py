import dataclasses
import zipfile

import pymupdf
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from brand import BrandError, load_brand
from deck import Deck
from deck_check import check_presentation


@pytest.fixture
def template_brand(make_brand):
    """Brand on python-pptx's default template (saved with 2 slides); `layouts` is the [deck.layouts] body."""
    prs = Presentation()
    for i in (0, 1):
        prs.slides.add_slide(prs.slide_layouts[i]).shapes.title.text = f"old {i}"

    def _make(layouts='title = "Title Slide"\ncontent = "Title Only"'):
        folder = make_brand(
            "tpl",
            f'[fonts]\nbody = "Calibri"\n[deck]\ntemplate = "template.pptx"\n'
            f'footer = "Lab X"\n[deck.layouts]\n{layouts}\n',
        )
        prs.save(str(folder / "template.pptx"))
        return folder

    return _make


def _empty_placeholders(slide):
    return [ph.name for ph in slide.placeholders if ph.has_text_frame and not ph.text_frame.text.strip()]


def test_template_slides_removed(template_brand, tmp_path):
    d = Deck(template_brand())
    assert len(d.prs.slides) == 0
    assert (d.W, d.H) == (10.0, 7.5)
    out = tmp_path / "out.pptx"
    assert d.save(out) == 0
    assert len(Presentation(str(out)).slides) == 0
    assert not [n for n in zipfile.ZipFile(out).namelist() if n.startswith("ppt/slides/")]


def test_template_slides_removed_then_added(template_brand, tmp_path):
    d = Deck(template_brand())
    d.content_slide("K", "Nuevo")
    out = tmp_path / "out.pptx"
    d.save(out)
    names = [n for n in zipfile.ZipFile(out).namelist() if n.startswith("ppt/slides/slide")]
    assert names == ["ppt/slides/slide1.xml"]
    assert [s.shapes.title.text for s in Presentation(str(out)).slides] == ["Nuevo"]


def test_title_uses_layout(template_brand):
    d = Deck(template_brand())
    slide = d.title_slide("k", "T", "sub")
    assert slide.slide_layout.name == "Title Slide"
    assert slide.shapes.title.text == "T"
    assert slide.shapes.title.name == "title"
    assert [ph.text_frame.text for ph in slide.placeholders] == ["T", "sub"]
    assert d.page == 1


def test_no_empty_placeholders(template_brand):
    d = Deck(template_brand('content = "Title and Content"'))
    s, top = d.content_slide("k", "T")
    assert s.slide_layout.name == "Title and Content"
    assert _empty_placeholders(s) == []
    assert [ph.placeholder_format.idx for ph in s.placeholders] == [0]
    t = s.shapes.title
    assert t.name == "title"
    assert top == pytest.approx((t.top + t.height) / 914400 + 0.25)


def test_section_and_final_use_layouts(template_brand):
    d = Deck(template_brand('section = "Section Header"\nfinal = "Title Slide"'))
    s = d.section_slide("01", "Sección")
    assert s.slide_layout.name == "Section Header"
    assert s.shapes.title.text == "Sección"
    assert _empty_placeholders(s) == []
    f = d.final_slide("Gracias", "")
    assert f.slide_layout.name == "Title Slide"
    assert [ph.text_frame.text for ph in f.placeholders] == ["Gracias"]


def test_section_layout_with_image(template_brand, tmp_path):
    png = tmp_path / "x.png"
    pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 30), False).save(str(png))
    d = Deck(template_brand('section = "Section Header"'))
    s = d.section_slide("02", "Con imagen", image=png, caption="Epígrafe")
    assert s.slide_layout.name == "Section Header"
    types = [sh.shape_type for sh in s.shapes]
    assert types.count(MSO_SHAPE_TYPE.PICTURE) == 1
    assert _empty_placeholders(s) == []
    # the picture goes behind the layout's text, so it never hides the title or the kicker
    assert types.index(MSO_SHAPE_TYPE.PICTURE) < min(i for i, sh in enumerate(s.shapes) if sh.is_placeholder)
    assert [i for i in check_presentation(d.prs) if i.kind == "covered"] == []


def test_missing_layout_lists_available(template_brand):
    d = Deck(template_brand('content = "Nope"'))
    with pytest.raises(BrandError, match="Nope") as e:
        d.content_slide("k", "T")
    assert "Title Only" in str(e.value)
    with pytest.raises(BrandError):
        d.layout("content")


def test_layout_unmapped_is_none(template_brand):
    d = Deck(template_brand())
    assert d.layout("final") is None
    assert d.layout("title").name == "Title Slide"


def test_unmapped_role_drawn(template_brand):
    d = Deck(template_brand())
    s = d.final_slide("Gracias", "Preguntas")
    assert s.slide_layout.name == "Blank"
    assert len(s.placeholders) == 0
    assert "Gracias" in [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]


def test_blank_slide_without_blank_layout_uses_fewest_placeholders(template_brand, tmp_path):
    d = Deck(template_brand())
    blank = next(l for l in d.prs.slide_layouts if l.name == "Blank")
    blank.name = "Vacío"
    s = d.blank_slide()
    assert s.slide_layout.name == "Vacío"
    assert len(s.placeholders) == 0


def test_blank_slide_removes_placeholders(template_brand):
    d = Deck(template_brand())
    d.prs.slide_layouts.remove(next(l for l in d.prs.slide_layouts if l.name == "Blank"))
    s = d.blank_slide()
    assert s.slide_layout.name == "Title Only"
    assert len(s.placeholders) == 0


def test_missing_template_file(template_brand, tmp_path):
    brand = load_brand(template_brand())
    gone = tmp_path / "gone.pptx"
    with pytest.raises(BrandError, match="gone.pptx"):
        Deck(dataclasses.replace(brand, template=gone))


def test_layouts_ignored_without_template(make_brand):
    folder = make_brand("plain", '[fonts]\nbody = "Calibri"\n[deck.layouts]\ntitle = "Title Slide"\n')
    d = Deck(folder)
    assert d.layout("title") is None
    assert d.title_slide("k", "T", "s").slide_layout.name == "Blank"


def _as_body(layout, top=None, height=None):
    """Turn the layout's OBJECT placeholder (idx 1) into a BODY one, optionally moved/resized (inches)."""
    ph = next(p for p in layout.placeholders if p.placeholder_format.idx == 1)
    ph._element.ph.set("type", "body")
    if top is not None:
        ph.top, ph.height = int(top * 914400), int(height * 914400)
    return ph


def test_template_blank_slide_inherits_background(template_brand):
    d = Deck(template_brand('title = "Title Slide"'))
    assert d.blank_slide()._element.cSld.bg is None
    assert d.content_slide("k", "Dibujada")[0]._element.cSld.bg is None
    assert d.blank_slide(bg="primary")._element.cSld.bg is not None
    assert d.final_slide("Gracias", "")._element.cSld.bg is not None


def test_mapped_content_slide_has_no_drawn_footer(template_brand):
    d = Deck(template_brand())
    s, _ = d.content_slide("k", "T")
    assert [sh.is_placeholder for sh in s.shapes] == [True]
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    assert "Lab X" not in texts and "01" not in texts


def test_kicker_not_written_into_large_body(template_brand):
    d = Deck(template_brand('content = "Title and Content"'))
    _as_body(next(l for l in d.prs.slide_layouts if l.name == "Title and Content"))
    s, top = d.content_slide("k", "T")
    assert [ph.placeholder_format.idx for ph in s.placeholders] == [0]
    assert "k" not in [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    t = s.shapes.title
    assert top == pytest.approx((t.top + t.height) / 914400 + 0.25)


def test_kicker_in_small_body_above_title(template_brand):
    d = Deck(template_brand('content = "Title and Content"'))
    layout = next(l for l in d.prs.slide_layouts if l.name == "Title and Content")
    title_top = layout.placeholders[0].top / 914400
    _as_body(layout, top=title_top - 0.5, height=0.4)
    s, _ = d.content_slide("k", "T")
    assert [ph.text_frame.text for ph in s.placeholders] == ["T", "k"]


def test_layout_without_title_adds_no_slide(template_brand):
    d = Deck(template_brand('content = "Blank"'))
    with pytest.raises(BrandError, match="no title placeholder"):
        d.content_slide("k", "T")
    assert len(d.prs.slides) == 0 and d.page == 0


def test_section_kicker_on_layout(template_brand):
    d = Deck(template_brand('section = "Section Header"'))
    s = d.section_slide("03", "Sección", kicker="PARTE")
    assert [ph.text_frame.text for ph in s.placeholders] == ["Sección", "PARTE 03"]
