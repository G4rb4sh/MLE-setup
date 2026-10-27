import pymupdf
import pytest
from pptx import Presentation
from pptx.enum.dml import MSO_FILL

from deck import Deck

STYLES = ("modern", "outlined", "dark")
ACCENT = "ABCDEF"
BORDER = "123456"


def _png(path, w, h):
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, w, h), False)
    pix.save(str(path))
    return path


def _deck(make_brand, style, footer="Lab X", **kw):
    folder = make_brand(
        f"b-{style}",
        f'style = "{style}"\n[colors]\naccent = "{ACCENT}"\ncard_border = "{BORDER}"\n'
        f'[fonts]\nbody = "Calibri"\n[deck]\nfooter = "{footer}"\n',
    )
    return Deck(folder, **kw)


def build_smoke(d, png):
    """One slide per component plus every slide factory; returns the number of slides built."""
    d.title_slide("Kicker", "Un título de portada", "Subtítulo de la presentación")
    d.section_slide("01", "Sección con número")
    d.section_slide("02", "Sección con imagen", image=png, caption="Epígrafe")

    s, top = d.content_slide("Cards", "Tres tarjetas")
    d.cards_row(s, top, [
        {"title": "Uno", "lines": ["Primera línea", "Segunda línea"], "tag": "etiqueta"},
        {"title": "Dos", "lines": ["Otra línea"], "accent": "highlight", "icon": png},
        {"title": "Tres", "lines": ["Una más", "Y otra"]},
    ])
    s = d.blank_slide()
    top = d.header(s, "Stats", "Números", lead="Una bajada corta bajo el título")
    d.footer(s)
    d.stat_row(s, top, [{"big": "93%", "cap": "precisión"}, {"big": "4x", "cap": "más rápido"}])
    s, top = d.content_slide("Compare", "Antes y después")
    d.compare(s, top, "Antes", ["Manual", "Lento"], "Después", ["Automático", "Rápido"])
    s, top = d.content_slide("Timeline", "Pasos")
    d.timeline(s, top + 0.4, [{"t": "Datos", "a": "Semana 1"}, {"t": "Modelo", "a": "Semana 2"},
                              {"t": "Deploy", "a": "Semana 3"}])
    s, top = d.content_slide("Chain", "Cadena")
    d.chain(s, d.col(0, 12)[0], top, d.col(0, 12)[1], 1.1,
            [{"t": "Ingesta", "a": "blob"}, {"t": "Índice"}, {"t": "Agente", "a": "loop"}])
    s, top = d.content_slide("Map", "Mapa")
    d.maprow(s, top, [("Chip", "Destacado", "resto"), ("Otro", "Texto", "")])
    s, top = d.content_slide("Grid", "Numerada")
    d.numbered_grid(s, top, [("1", "Uno"), ("2", "Dos"), ("3", "Tres"), ("4", "Cuatro")])
    s, top = d.content_slide("Code", "Código")
    d.code_block(s, top, ["# comentario", "x = 1", "print(x)"], note="Una nota")
    s, top = d.content_slide("Images", "Fila de imágenes")
    d.image_row(s, top, [{"path": png, "caption": "A"}, {"path": png, "caption": "B"}])
    s, top = d.content_slide("Split", "Imagen y bullets")
    d.split_image(s, top, png, ["Uno: detalle", "Dos"], caption="Epígrafe")
    s, top = d.content_slide("Diagram", "Diagrama")
    d.container(s, 0.6, top, 12.1, 3.2, "PIPELINE")
    d.node(s, 1.0, top + 0.8, 2.5, 1.2, "Hecho", "sub")
    d.labeled_arrow(s, 3.6, top + 1.3, 1.6, "pasa")
    d.node(s, 5.3, top + 0.8, 2.5, 1.2, "Planeado", state="planned")
    d.node(s, 8.4, top + 0.8, 2.5, 1.2, "Recto", shape="rect")
    d.legend(s, 0.6, 6.4)
    d.image_slide("Image", "Imagen grande", png, caption="Epígrafe")
    d.final_slide("Gracias", "Preguntas")
    return 16


@pytest.mark.parametrize("style", STYLES)
def test_smoke_all_components(make_brand, tmp_path, style):
    d = _deck(make_brand, style)
    n = build_smoke(d, _png(tmp_path / "img.png", 300, 200))
    out = tmp_path / f"{style}.pptx"
    assert d.save(out) == n
    prs = Presentation(str(out))
    assert len(prs.slides) == n
    W, H = prs.slide_width, prs.slide_height
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            assert 0 <= sh.left and sh.left + sh.width <= W, (i, sh.name)
            assert 0 <= sh.top and sh.top + sh.height <= H, (i, sh.name)


def _card(slide):
    return next(sh for sh in slide.shapes if sh.name == "card")


def test_card_style_differs(make_brand):
    lines = {}
    for style in STYLES:
        d = _deck(make_brand, style)
        s, top = d.content_slide("k", "t")
        d.cards_row(s, top, [{"title": "A", "lines": ["x"]}])
        lines[style] = _card(s).line
    assert lines["outlined"].color.rgb == d.color(ACCENT)
    assert lines["outlined"].width.pt == pytest.approx(1.5)
    assert lines["modern"].color.rgb == d.color(BORDER)
    assert lines["dark"].fill.type in (MSO_FILL.BACKGROUND, None)


def test_card_accent_strip_follows_components(make_brand):
    for style, strip in (("modern", "top"), ("outlined", "left")):
        d = _deck(make_brand, style)
        s, top = d.content_slide("k", "t")
        d.cards_row(s, top, [{"title": "A", "lines": ["x"], "accent": "C0FFEE"}])
        card = _card(s)
        strips = [sh for sh in s.shapes if sh.shape_type == 1 and sh.name != "card"
                  and sh.fill.type == MSO_FILL.SOLID and sh.fill.fore_color.rgb == d.color("C0FFEE")]
        assert len(strips) == 1
        st = strips[0]
        if strip == "top":
            assert st.width == card.width and st.height < card.height / 4
        else:
            assert st.height == card.height and st.width < card.width / 4


def test_planned_node_dashed(make_brand):
    d = _deck(make_brand, "dark")
    s = d.blank_slide()
    done = d.node(s, 1, 1, 2, 1, "A")
    planned = d.node(s, 4, 1, 2, 1, "B", state="planned")
    assert done.line.dash_style is None
    assert planned.line.dash_style is not None
    assert planned.line.color.rgb == d.color("planned")
    with pytest.raises(ValueError, match="state"):
        d.node(s, 1, 3, 2, 1, "C", state="maybe")


def test_content_slide_returns_top(make_brand):
    d = _deck(make_brand, "modern")
    s, top = d.content_slide("Kicker", "El título")
    assert 1.0 < top < 2.5
    titles = [sh for sh in s.shapes if sh.name == "title"]
    assert len(titles) == 1 and titles[0].text_frame.text == "El título"


def test_header_lead_pushes_top_down(make_brand):
    d = _deck(make_brand, "modern")
    plain = d.header(d.blank_slide(), "k", "t")
    with_lead = d.header(d.blank_slide(), "k", "t", lead="Una bajada")
    assert with_lead > plain


def _all_text(slide):
    return " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def test_footer_from_brand(make_brand):
    d = _deck(make_brand, "modern")
    s, _ = d.content_slide("k", "t")
    text = _all_text(s)
    assert "Lab X" in text
    assert "01" in text


def test_footer_override(make_brand):
    d = _deck(make_brand, "modern", footer="Lab X")
    d2 = Deck(d.brand, footer="Otro pie")
    s, _ = d2.content_slide("k", "t")
    text = _all_text(s)
    assert "Otro pie" in text and "Lab X" not in text


def test_dark_heading_readable(make_brand):
    d = _deck(make_brand, "dark")
    s, _ = d.content_slide("k", "Título")
    title = next(sh for sh in s.shapes if sh.name == "title")
    assert title.text_frame.paragraphs[0].runs[0].font.color.rgb == d.color("on_primary")
    d = _deck(make_brand, "modern")
    s, _ = d.content_slide("k", "Título")
    title = next(sh for sh in s.shapes if sh.name == "title")
    assert title.text_frame.paragraphs[0].runs[0].font.color.rgb == d.color("primary")


def test_image_row_webp_message(make_brand, tmp_path):
    d = _deck(make_brand, "modern")
    p = tmp_path / "x.webp"
    p.write_bytes(b"RIFF")
    with pytest.raises(ValueError, match=r"\.webp no soportado"):
        d.image_row(d.blank_slide(), 2, [{"path": p}])
