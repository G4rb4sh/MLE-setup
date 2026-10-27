import datetime as dt
import os
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import docx_brand
import md2docx
from brand import load_brand
from md2docx import DocMeta, build_docx, parse_front_matter, render_mermaid
from render import applescript_quote
from test_components import _png

pytestmark = pytest.mark.skipif(shutil.which("pandoc") is None, reason="needs pandoc")

DATA = Path(__file__).resolve().parent / "data"
TITLE = "Plan de producción"


def _copy_sample(folder: Path) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    md = folder / "sample.md"
    shutil.copyfile(DATA / "sample.md", md)
    _png(folder / "img.png", 120, 80)
    return md


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    md = _copy_sample(tmp_path_factory.mktemp("md") / "Docs con espacio — ñ")
    out = build_docx(md, load_brand("neutral"), use_word=False)
    return out


def _text(doc) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


def _rels(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read("word/_rels/document.xml.rels").decode()


def test_front_matter_parse():
    meta_d, body = parse_front_matter((DATA / "sample.md").read_text(encoding="utf-8"))
    assert meta_d == {"title": TITLE, "subtitle": "Documento de ejemplo para las pruebas",
                      "client": "Demo", "version": "0.9"}
    assert body.lstrip().startswith("# Plan de producción")
    meta = DocMeta.from_dict(meta_d, fallback_title="sample")
    assert (meta.title, meta.client, meta.version) == (TITLE, "Demo", "0.9")
    with pytest.raises(ValueError, match="version"):
        DocMeta.from_dict({"colour": "x"}, fallback_title="t")


def test_front_matter_absent_and_quoted():
    assert parse_front_matter("# Hola\n\ntexto\n") == ({}, "# Hola\n\ntexto\n")
    assert parse_front_matter("---\nsolo una regla\n") == ({}, "---\nsolo una regla\n")
    d, body = parse_front_matter('---\ntitle: "A: b"\r\nauthor: \'Ana\'\n---\ncuerpo\n')
    assert d == {"title": "A: b", "author": "Ana"} and body == "cuerpo\n"
    with pytest.raises(ValueError, match="key: value"):
        parse_front_matter("---\nsin dos puntos\n---\n")


def test_docmeta_defaults_and_types():
    m = DocMeta.from_dict({}, fallback_title="stem")
    today = dt.date.today()
    assert m.title == "stem" and m.version == "1.0" and m.lang == "es-UY"
    assert m.date.startswith(f"{today.day} de ") and m.date.endswith(str(today.year))
    assert (m.toc_depth, m.page_break_on_h1, m.shift_headings) == (3, True, 0)
    m = DocMeta.from_dict({"toc_depth": "2", "page_break_on_h1": "no", "shift_headings": "-1",
                           "date": "1 de mayo"}, fallback_title="t")
    assert (m.toc_depth, m.page_break_on_h1, m.shift_headings, m.date) == (2, False, -1, "1 de mayo")
    with pytest.raises(ValueError, match="toc_depth"):
        DocMeta.from_dict({"toc_depth": "tres"}, fallback_title="t")
    with pytest.raises(ValueError, match="page_break_on_h1"):
        DocMeta.from_dict({"page_break_on_h1": "quizás"}, fallback_title="t")


def test_docx_built(built):
    assert built.exists() and built.name == "sample.docx"
    doc = Document(str(built))
    text = _text(doc)
    assert TITLE in text and "0.9" in text
    assert len(doc.tables) >= 1
    assert len(doc.inline_shapes) >= 1


def test_cross_link_italic(built):
    doc = Document(str(built))
    runs = [r for p in doc.paragraphs for r in p.runs if r.text == "Arquitectura"]
    assert runs and all(r.italic is True for r in runs)
    assert "arquitectura.md" not in _rels(built)


def test_heading_font_from_brand(built):
    doc = Document(str(built))
    assert doc.styles["Heading 1"].font.name == load_brand("neutral").fonts["heading"]


def test_header_has_title(built):
    doc = Document(str(built))
    sec = doc.sections[0]
    assert TITLE in "\n".join(p.text for p in sec.header.paragraphs)
    assert sec.different_first_page_header_footer  # the cover page has no header
    assert "PAGE" in sec.footer._element.xml


def test_cover_toc_and_body_order(built):
    doc = Document(str(built))
    body = list(doc.element.body)
    toc = doc.element.body.find(qn("w:sdt"))
    assert toc is not None
    first_heading = next(p for p in doc.paragraphs if p.style.name == "Heading 1")
    assert first_heading.text == "Contexto"  # the leading H1 equal to the title moved to the cover
    cover_idx = next(i for i, el in enumerate(body) if TITLE in "".join(el.itertext()))
    assert cover_idx < body.index(toc) < body.index(first_heading._p)
    assert "w:anchor" in body[0].xml or any("wp:anchor" in el.xml for el in body[:2])
    # The front-matter `---` rule and the body `---` are not drawn as lines.
    assert not any(p.text.strip() == "---" for p in doc.paragraphs)


def test_tables_styled_from_brand(built):
    brand = load_brand("neutral")
    doc = Document(str(built))
    tbl = doc.tables[0]
    assert len(tbl.rows) == 3 and len(tbl.columns) == 2
    xml = tbl._tbl.xml
    assert f'w:color="{brand.colors["accent"]}"' in xml  # heading_color falls back to accent
    assert f'w:fill="{brand.colors["accent"]}"' in tbl.rows[0].cells[0]._tc.xml
    assert "w:tblHeader" in xml


def test_update_fields_without_word(built):
    doc = Document(str(built))
    assert doc.settings.element.find(qn("w:updateFields")) is not None


def test_styles_and_colors_from_brand(make_brand, tmp_path):
    folder = make_brand("b", '[colors]\naccent = "112233"\ntext = "445566"\nbackground = "FAFAFA"\n'
                             'card_border = "CCCCCC"\n'
                             '[fonts]\nbody = "Verdana"\nheading = "Georgia"\nmono = "Courier New"\n')
    brand = load_brand(folder)
    ref = docx_brand.make_reference_docx(brand, tmp_path / "ref.docx", page_break_on_h1=False)
    s = Document(str(ref)).styles
    assert s["Heading 1"].font.name == "Georgia"
    assert str(s["Heading 1"].font.color.rgb) == "112233"
    assert not s["Heading 1"].paragraph_format.page_break_before
    assert s["Normal"].font.name == "Verdana" and str(s["Normal"].font.color.rgb) == "445566"
    assert s["Source Code"].font.name == "Courier New"
    assert 'w:fill="F2F2F2"' in s["Source Code"].element.xml
    assert 'w:fill="FAFAFA"' in s["Block Text"].element.xml

    folder = make_brand("c", '[fonts]\nbody = "Verdana"\n[docx]\nheading_color = "highlight"\n'
                             'code_bg = "EEEEEE"\nnote_bg = "#abcdef"\nrule = "010203"\n')
    brand = load_brand(folder)
    ref = docx_brand.make_reference_docx(brand, tmp_path / "ref2.docx", page_break_on_h1=True)
    s = Document(str(ref)).styles
    assert str(s["Heading 2"].font.color.rgb) == brand.colors["highlight"]
    assert s["Heading 1"].paragraph_format.page_break_before
    assert 'w:fill="EEEEEE"' in s["Source Code"].element.xml
    assert 'w:fill="ABCDEF"' in s["Block Text"].element.xml


def test_brand_reference_docx_used(make_brand, tmp_path):
    folder = make_brand("r", '[fonts]\nbody = "Verdana"\n[docx]\nreference = "ref.docx"\n')
    d = Document()
    d.styles["Heading 1"].font.name = "Brand Font"
    d.save(str(folder / "ref.docx"))
    ref = docx_brand.make_reference_docx(load_brand(folder), tmp_path / "out.docx", True)
    s = Document(str(ref)).styles
    assert s["Heading 1"].font.name == "Brand Font"
    assert s["Heading 1"].paragraph_format.page_break_before


def test_cover_band_logo_and_rule(make_brand, tmp_path):
    folder = make_brand("l", '[colors]\nprimary = "102030"\ncard_border = "ABABAB"\n'
                             '[fonts]\nbody = "Verdana"\n[logos]\nlight = "l.png"\ndark = "d.png"\n')
    _png(folder / "l.png", 200, 50)
    _png(folder / "d.png", 200, 50)
    md = _copy_sample(tmp_path / "doc")
    out = build_docx(md, load_brand(folder), out=tmp_path / "x.docx", use_word=False)
    doc = Document(str(out))
    body_xml = doc.element.body.xml
    assert '<a:srgbClr val="102030"/>' in body_xml  # cover band = primary
    assert body_xml.count("<pic:pic") >= 2  # cover logo + markdown image
    sec = doc.sections[0]
    assert sec.header._element.xml.count("<pic:pic") == 1
    assert 'w:color="ABABAB"' in sec.footer._element.xml  # rule falls back to card_border


def test_mermaid_missing_tools(monkeypatch, tmp_path, capsys):
    which = shutil.which
    monkeypatch.setattr(md2docx.shutil, "which", lambda n, *a, **k: None if n == "npx" else which(n))
    md = tmp_path / "m.md"
    md.write_text("# Diagrama\n\n```mermaid\nflowchart LR\n  A --> B\n```\n\n"
                  "Texto.\n\n```mermaid\nflowchart LR\n  C --> D\n```\n", encoding="utf-8")
    out = build_docx(md, load_brand("neutral"), use_word=False)
    text = _text(Document(str(out)))
    assert "A --> B" in text and "C --> D" in text
    err = capsys.readouterr().err
    assert err.count("mermaid") == 1 and "npx" in err


def test_render_mermaid_lr_to_tb(monkeypatch, tmp_path):
    calls = []

    def fake_run(cmd, **kw):
        src = Path(cmd[cmd.index("-i") + 1])
        calls.append(src.read_text(encoding="utf-8"))
        _png(Path(cmd[cmd.index("-o") + 1]), 40, 30)
        return subprocess.CompletedProcess(cmd, 0, "", "")

    monkeypatch.setattr(md2docx, "_chrome", lambda: "/fake/chrome")
    monkeypatch.setattr(md2docx.shutil, "which", lambda n, *a, **k: "/fake/npx")
    monkeypatch.setattr(md2docx.subprocess, "run", fake_run)
    md = "Antes\n\n```mermaid\nflowchart LR\n  A --> B\n```\n\nDespués\n"
    out = render_mermaid(md, tmp_path)
    assert calls == ["flowchart TB\n  A --> B\n"]
    assert "```mermaid" not in out and "Antes" in out and "Después" in out
    png = re.search(r"!\[\]\(<(.+?)>\)", out).group(1)
    assert Path(png).exists() and Path(png).parent == tmp_path


def test_render_mermaid_failure_keeps_block(monkeypatch, tmp_path, capsys):
    def failing(cmd, **kw):
        raise subprocess.CalledProcessError(1, cmd, "", "Parse error")

    monkeypatch.setattr(md2docx, "_chrome", lambda: "/fake/chrome")
    monkeypatch.setattr(md2docx.shutil, "which", lambda n, *a, **k: "/fake/npx")
    monkeypatch.setattr(md2docx.subprocess, "run", failing)
    md = "```mermaid\nflowchart LR\n  A --> B\n```\n"
    assert render_mermaid(md, tmp_path) == md
    assert "Parse error" in capsys.readouterr().err


def test_use_word_skips_update_on_open(monkeypatch, tmp_path):
    seen = []
    monkeypatch.setattr(md2docx, "fill_toc_with_word", lambda p: seen.append(p) or True)
    md = _copy_sample(tmp_path / "w")
    out = build_docx(md, load_brand("neutral"), out=tmp_path / "w.docx")
    assert seen == [out]
    assert Document(str(out)).settings.element.find(qn("w:updateFields")) is None


def test_cli(tmp_path, capsys):
    md = _copy_sample(tmp_path / "cli")
    out = tmp_path / "salida.docx"
    assert md2docx.main([str(md), "--brand", "neutral", "-o", str(out), "--no-word"]) == 0
    assert out.exists() and str(out) in capsys.readouterr().out
    assert md2docx.main([str(md), "--brand", "nope-brand"]) == 2
    assert "nope-brand" in capsys.readouterr().err
    bad = tmp_path / "bad.md"
    bad.write_text("---\ncolour: x\n---\n# T\n", encoding="utf-8")
    assert md2docx.main([str(bad), "--brand", "neutral", "--no-word"]) == 2
    assert "version" in capsys.readouterr().err


def test_fix_property_order(tmp_path):
    path = tmp_path / "o.docx"
    d = Document()
    p = d.add_paragraph("x")
    ppr = p._p.get_or_add_pPr()
    ppr.append(OxmlElement("w:jc"))
    ppr.append(OxmlElement("w:pBdr"))
    ppr.insert(0, OxmlElement("w:spacing"))
    d.save(str(path))
    docx_brand.fix_property_order(path)
    ppr = Document(str(path)).paragraphs[0]._p.pPr
    assert [c.tag.split("}")[1] for c in ppr] == ["pBdr", "spacing", "jc"]


# --------------------------------------------------------------------------- #
# Word TOC: only ever the temporary copy, addressed by name.

def test_toc_script_targets_copy_by_name():
    copy = Path('/var/x/Carpeta "q" — ñ/office-docs-0a1b2c3d.docx')
    script = docx_brand._toc_script(copy)
    assert f"set nm to {applescript_quote(Path(copy.name))}" in script
    acts = [ln.strip() for ln in script.splitlines() if re.match(r"\s*(close|save|update)\b", ln)]
    assert len(acts) == 4  # update, save, close, best-effort close on the late-open path
    assert all("document nm" in ln for ln in acts)
    assert "active document" not in script
    assert not re.search(r"document (idx|i|\d+)\b", script)


def _docx(path: Path) -> Path:
    d = Document()
    d.add_paragraph("original")
    d.save(str(path))
    return path


def test_fill_toc_without_word(monkeypatch, tmp_path):
    monkeypatch.setattr(docx_brand, "office_app", lambda suffix: None)
    assert docx_brand.fill_toc_with_word(_docx(tmp_path / "a.docx")) is False


@pytest.mark.parametrize("err", ["Not authorized to send Apple events to Microsoft Word. (-1743)",
                                 "table of contents 1 doesn't exist (-1728)"])
def test_fill_toc_failure_leaves_file(monkeypatch, tmp_path, capsys, err):
    monkeypatch.setattr(docx_brand, "office_app", lambda suffix: "Microsoft Word")
    monkeypatch.setattr(docx_brand.subprocess, "run",
                        lambda cmd, **kw: subprocess.CompletedProcess(cmd, 1, "", err))
    path = _docx(tmp_path / "a.docx")
    before = path.read_bytes()
    assert docx_brand.fill_toc_with_word(path) is False
    assert path.read_bytes() == before
    msg = capsys.readouterr().err
    assert ("Privacy & Security" in msg) == ("-1743" in err)


def test_fill_toc_timeout_cleans_up_by_name(monkeypatch, tmp_path):
    scripts = []

    def run(cmd, input, **kw):
        scripts.append(input)
        if len(scripts) == 1:
            raise subprocess.TimeoutExpired(cmd, 1)
        return subprocess.CompletedProcess(cmd, 0, "", "")

    monkeypatch.setattr(docx_brand, "office_app", lambda suffix: "Microsoft Word")
    monkeypatch.setattr(docx_brand.subprocess, "run", run)
    assert docx_brand.fill_toc_with_word(_docx(tmp_path / "a.docx")) is False
    name = re.search(r'set nm to "(office-docs-[0-9a-f]{8}\.docx)"', scripts[0]).group(1)
    assert len(scripts) == 2 and f'close document "{name}" saving no' in scripts[1]


def test_fill_toc_copies_result_back(monkeypatch, tmp_path):
    def run(cmd, input, **kw):
        src = Path(re.search(r'set src to "(.+)"', input).group(1))
        assert src.name.startswith("office-docs-") and src.parent != tmp_path
        src.write_bytes(b"updated by word")
        return subprocess.CompletedProcess(cmd, 0, "", "")

    monkeypatch.setattr(docx_brand, "office_app", lambda suffix: "Microsoft Word")
    monkeypatch.setattr(docx_brand.subprocess, "run", run)
    path = _docx(tmp_path / "a.docx")
    assert docx_brand.fill_toc_with_word(path) is True
    assert path.read_bytes() == b"updated by word"


def _word_documents() -> list[str]:
    r = subprocess.run(["osascript", "-e", 'tell application "Microsoft Word" to get name of documents'],
                       capture_output=True, text=True, check=True)
    return sorted(n.strip() for n in r.stdout.strip().split(",") if n.strip())


@pytest.mark.skipif(os.environ.get("OFFICE_DOCS_RENDER") != "1", reason="set OFFICE_DOCS_RENDER=1")
def test_real_word_toc(tmp_path):
    md = _copy_sample(tmp_path / "Docs con espacio — ñ")
    before = _word_documents()
    out = build_docx(md, load_brand("neutral"))
    assert _word_documents() == before
    doc = Document(str(out))
    toc = doc.element.body.find(qn("w:sdt"))
    assert "Contexto" in "".join(toc.itertext())
    assert doc.settings.element.find(qn("w:updateFields")) is None
