import pymupdf
import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from brand import load_brand
from deck import Deck


@pytest.fixture
def d(tmp_path):
    return Deck("neutral", cwd=tmp_path)


def _png(path, w, h):
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, w, h), False)
    pix.save(str(path))
    return path


def test_deck_size_and_brand_instance(tmp_path):
    brand = load_brand("neutral", cwd=tmp_path)
    d = Deck(brand, footer="Custom")
    assert d.brand is brand
    assert round(d.W, 3) == 13.333 and d.H == 7.5
    assert d.page == 0
    assert d.footer_text == "Custom"


def test_color_tokens_and_hex(d):
    assert d.color("accent") == RGBColor.from_string(d.brand.colors["accent"])
    assert d.color("#00ff00") == RGBColor(0, 0xFF, 0)
    with pytest.raises(ValueError, match="navy") as e:
        d.color("navy")
    assert "accent" in str(e.value)


def test_font_roles_and_literal(d):
    assert d.font("mono") == "Consolas"
    assert d.font("Georgia") == "Georgia"


def test_col_grid(d):
    x, w = d.col(0, 12)
    assert x == pytest.approx(0.6)
    assert w == pytest.approx(d.W - 1.2)
    x1, w1 = d.col(1)
    assert x1 == pytest.approx(0.6 + w1 + 0.2)
    with pytest.raises(ValueError):
        d.col(11, 2)


def test_blank_slide_bg_and_page(d):
    s = d.blank_slide()
    assert d.page == 1
    assert s.background.fill.fore_color.rgb == d.color(d.brand.components["content_bg"])
    s2 = d.blank_slide(bg="primary")
    assert s2.background.fill.fore_color.rgb == d.color("primary")


def test_rect_units_and_fill(d):
    s = d.blank_slide()
    sh = d.rect(s, 1, 1, 2, 1, fill="primary")
    assert sh.left == Inches(1)
    assert sh.width == Inches(2)
    assert sh.fill.fore_color.rgb == d.color("primary")


def test_rect_dash_and_line(d):
    s = d.blank_slide()
    sh = d.rect(s, 0, 0, 1, 1, fill=None, line="accent", dash=True, radius=0.1)
    assert sh.line.color.rgb == d.color("accent")
    assert sh.line.dash_style is not None


def test_line_arrow_end(d):
    s = d.blank_slide()
    ln = d.line(s, 0, 0, 2, 1, arrow_end=True)
    assert "tailEnd" in ln._element.xml


def test_arrow_direction(d):
    s = d.blank_slide()
    d.arrow(s, 0, 0, 1, 0.5, direction="down")
    with pytest.raises(ValueError, match="sideways"):
        d.arrow(s, 0, 0, 1, 0.5, direction="sideways")


def test_text_runs(d):
    s = d.blank_slide()
    sh = d.text(s, 0, 0, 4, 1, [[{"text": "A", "bold": True}, "b"], "second"])
    paras = sh.text_frame.paragraphs
    assert len(paras) == 2
    runs = paras[0].runs
    assert [r.text for r in runs] == ["A", "b"]
    assert runs[0].font.bold is True
    assert paras[1].text == "second"


def test_text_newline_splits_paragraphs(d):
    s = d.blank_slide()
    sh = d.text(s, 0, 0, 4, 1, "one\ntwo", color="accent", font="mono")
    paras = sh.text_frame.paragraphs
    assert [p.text for p in paras] == ["one", "two"]
    assert paras[0].runs[0].font.name == "Consolas"
    assert paras[0].runs[0].font.color.rgb == d.color("accent")


def test_bullets_lead_and_sub(d):
    s = d.blank_slide()
    sh = d.bullets(s, 0, 0, 6, 3, ["Uno: detalle", "  sub"])
    paras = sh.text_frame.paragraphs
    assert any(r.text == "Uno: " and r.font.bold for r in paras[0].runs)
    assert "sub" in paras[1].text
    lead = next(r for r in paras[0].runs if r.text == "Uno: ")
    assert lead.font.color.rgb == d.color("primary")
    sh = d.bullets(s, 0, 0, 6, 3, ["Uno: detalle"], lead_color="accent")
    lead = next(r for r in sh.text_frame.paragraphs[0].runs if r.text == "Uno: ")
    assert lead.font.color.rgb == d.color("accent")


def test_chip_autosize(d):
    s = d.blank_slide()
    short = d.chip(s, 0, 0, "AB")
    long = d.chip(s, 0, 1, "A much longer label")
    assert long.width > short.width
    assert long.text_frame.text == "A much longer label"


def test_image_webp_rejected(d, tmp_path):
    s = d.blank_slide()
    p = tmp_path / "x.webp"
    p.write_bytes(b"RIFF")
    with pytest.raises(ValueError, match=r"\.webp no soportado por python-pptx: convertir a PNG"):
        d.image(s, p, 0, 0, 2, 2)


def test_image_missing(d, tmp_path):
    s = d.blank_slide()
    p = tmp_path / "nope.png"
    with pytest.raises(FileNotFoundError, match="nope.png"):
        d.image(s, p, 0, 0, 2, 2)


def test_image_contain_fit(d, tmp_path):
    s = d.blank_slide()
    p = _png(tmp_path / "wide.png", 200, 100)
    pic = d.image(s, p, 1, 1, 2, 2)
    assert pic.width == Inches(2)
    assert pic.height == Inches(1)
    assert pic.top == Inches(1.5)
    assert pic.left == Inches(1)


def test_image_caption_color(d, tmp_path):
    s = d.blank_slide()
    p = _png(tmp_path / "c.png", 20, 20)
    d.image(s, p, 0, 0, 2, 2, caption="def")
    d.image(s, p, 3, 0, 2, 2, caption="own", caption_color="on_primary")
    runs = {sh.text_frame.text: sh.text_frame.paragraphs[0].runs[0] for sh in s.shapes if sh.has_text_frame}
    assert runs["def"].font.color.rgb == d.color("muted")
    assert runs["own"].font.color.rgb == d.color("on_primary")


def test_icon_square_box(d, tmp_path):
    s = d.blank_slide()
    p = _png(tmp_path / "sq.png", 50, 50)
    pic = d.icon(s, p, 0, 0, size=0.5)
    assert pic.width == Inches(0.5) and pic.height == Inches(0.5)


def test_scrim_alpha(d):
    s = d.blank_slide()
    sh = d.scrim(s, 0, 0, 2, 2, "primary", 40)
    assert '<a:alpha val="40000"/>' in sh._element.xml


def test_gradient_alphas(d):
    s = d.blank_slide()
    sh = d.gradient(s, 0, 0, 2, 2, "primary", 80, 0, angle=90)
    xml = sh._element.xml
    assert "gradFill" in xml and "solidFill" not in xml.split("<a:ln")[0]
    assert '<a:alpha val="80000"/>' in xml and '<a:alpha val="0"/>' in xml
    assert 'ang="5400000"' in xml


def test_table_header(d):
    s = d.blank_slide()
    gf = d.table(s, 0, 0, 4, [["H1", "H2"], ["a", "b"]])
    cell = gf.table.cell(0, 0)
    assert cell.fill.fore_color.rgb == d.color("primary")
    assert cell.text_frame.paragraphs[0].runs[0].font.color.rgb == d.color("on_primary")


def test_notes(d):
    s = d.blank_slide()
    d.notes(s, "  hola  ")
    assert s.notes_slide.notes_text_frame.text == "hola"


def test_save_counts(d, tmp_path):
    d.blank_slide()
    d.blank_slide()
    out = tmp_path / "out.pptx"
    assert d.save(out) == 2
    assert len(Presentation(str(out)).slides) == 2
