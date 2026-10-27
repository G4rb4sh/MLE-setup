import os
import re
import struct
import tomllib
import zipfile
import zlib

import pytest
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches

from brand import PLUGIN_BRANDS, BrandError, load_brand
from brand_from_pptx import brand_from_pptx, confirm_brand
from deck import Deck
from new_brand import new_brand

LABEL_XML = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
    b'<clbl:labelList xmlns:clbl="http://schemas.microsoft.com/office/2020/mipLabelMetadata">'
    b'<clbl:label id="{00000000-0000-0000-0000-000000000001}" enabled="1" method="Standard" '
    b'siteId="{00000000-0000-0000-0000-000000000002}" removed="0"/></clbl:labelList>'
)
LABEL_REL = (
    '<Relationship Id="rIdLabel" Type="http://schemas.microsoft.com/office/2020/02/relationships/'
    'classificationlabels" Target="docMetadata/LabelInfo.xml"/>'
)
LABEL_CT = ('<Override PartName="/docMetadata/LabelInfo.xml" '
            'ContentType="application/vnd.ms-office.classificationlabels+xml"/>')
SECRET = "Proyecto Secreto"
CUSTOM_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="MSIP_Label_1_Enabled">'
    '<vt:lpwstr>true</vt:lpwstr></property>'
    '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="3" name="Cliente">'
    f'<vt:lpwstr>{SECRET}</vt:lpwstr></property></Properties>'
).encode()
CUSTOM_REL = ('<Relationship Id="rIdCustom" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
              'relationships/custom-properties" Target="docProps/custom.xml"/>')
CUSTOM_CT = ('<Override PartName="/docProps/custom.xml" '
             'ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/>')


def _png_bytes() -> bytes:
    """A valid 2x2 RGB PNG."""
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    raw = b"".join(b"\x00" + b"\x20\x70\xcf" * 2 for _ in range(2))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


_PNG = _png_bytes()


@pytest.fixture
def client_pptx(tmp_path):
    """A client-like deck: 2 slides (with notes) of SECRET, "Cliente X" on the Title Slide layout,
    author "Ana", a Company in app.xml, a sensitivity label (LabelInfo.xml and an MSIP_Label_ custom
    property) and a custom property holding SECRET injected into the zip."""
    prs = Presentation()
    slides = []
    for _ in range(2):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = SECRET
        s.notes_slide.notes_text_frame.text = f"{SECRET}: notas"
        slides.append(s)
    layout = next(lay for lay in prs.slide_layouts if lay.name == "Title Slide")
    tb = slides[0].shapes.add_textbox(Inches(1), Inches(6.5), Inches(3), Inches(0.5))
    tb.text_frame.text = "Cliente X"
    tb.name = "TextBox 3"
    tb._element.nvSpPr.cNvPr.id = 99
    layout.shapes._spTree.append(tb._element)  # moves the textbox from the slide to the layout
    prs.core_properties.author = "Ana"
    prs.core_properties.title = SECRET
    plain = tmp_path / "plain.pptx"
    prs.save(str(plain))

    src = tmp_path / "client.pptx"
    with zipfile.ZipFile(plain) as zin, zipfile.ZipFile(src, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", (LABEL_CT + CUSTOM_CT).encode() + b"</Types>")
            elif item.filename == "_rels/.rels":
                data = data.replace(b"</Relationships>", (LABEL_REL + CUSTOM_REL).encode() + b"</Relationships>")
            elif item.filename == "docProps/app.xml":
                data = re.sub(rb"<(Company|Manager)>.*?</\1>|<(Company|Manager)/>", b"", data)
                data = data.replace(b"</Properties>",
                                    f"<Company>{SECRET} SA</Company><Manager>Ana</Manager></Properties>"
                                    .encode())
            zout.writestr(item, data)
        zout.writestr("docMetadata/LabelInfo.xml", LABEL_XML)
        zout.writestr("docProps/custom.xml", CUSTOM_XML)
    return src


@pytest.fixture
def extracted(client_pptx, tmp_path):
    return brand_from_pptx(client_pptx, "cliente", tmp_path / "brands", render_samples=False)


def _parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def test_slides_removed_and_props_cleared(extracted):
    template = extracted.brand_dir / "template.pptx"
    prs = Presentation(str(template))
    assert len(prs.slides) == 0
    assert prs.core_properties.author == ""
    assert prs.core_properties.title == ""
    parts = _parts(template)
    assert [n for n in parts if SECRET.encode() in parts[n]] == []
    assert not [n for n in parts
                if re.match(r"ppt/(slides|notesSlides|comments|printerSettings)/|docProps/thumbnail", n)]
    assert b"<Company>" not in parts["docProps/app.xml"].replace(b"<Company></Company>", b"")
    assert b"Ana" not in parts["docProps/app.xml"]


def test_template_package_is_consistent(extracted):
    parts = _parts(extracted.brand_dir / "template.pptx")
    overrides = re.findall(rb'PartName="/([^"]+)"', parts["[Content_Types].xml"])
    assert overrides and all(o.decode() in parts for o in overrides)
    assert len(Presentation(str(extracted.brand_dir / "template.pptx")).slide_layouts) == 11


def test_layout_text_reported(extracted):
    assert any(text == "Cliente X" and "Title Slide" in loc for loc, text in extracted.texts)
    assert ("layout 'Title Slide' / TextBox 3", "Cliente X") in extracted.texts
    # placeholder prompts are listed too: they can carry client text as well
    assert any(text == "Click to edit Master title style" and loc.startswith("master")
               for loc, text in extracted.texts)
    assert "Title Slide" in extracted.layouts and "Title and Content" in extracted.layouts


def test_label_kept(extracted, client_pptx):
    assert extracted.has_label
    parts = _parts(extracted.brand_dir / "template.pptx")
    assert parts["docMetadata/LabelInfo.xml"] == LABEL_XML
    assert b"docMetadata/LabelInfo.xml" in parts["_rels/.rels"]
    assert b'name="MSIP_Label_1_Enabled"' in parts["docProps/custom.xml"]
    assert b'name="Cliente"' not in parts["docProps/custom.xml"]


def test_theme_mapped_to_draft_toml(extracted):
    data = tomllib.loads((extracted.brand_dir / "brand.toml").read_text(encoding="utf-8"))
    assert data["draft"] is True and data["style"] == "outlined"
    c = data["colors"]  # python-pptx's default "Office Theme"
    assert (c["primary"], c["accent"], c["highlight"]) == ("1F497D", "4F81BD", "C0504D")
    assert (c["text"], c["background"], c["surface"]) == ("000000", "FFFFFF", "EEECE1")
    assert c["on_primary"] == "FFFFFF"
    assert data["fonts"] == {"heading": "Calibri", "body": "Calibri"}
    assert data["deck"]["layouts"] == {"title": "Title Slide", "content": "Title and Content",
                                       "section": "Section Header"}
    assert "client.pptx" not in (extracted.brand_dir / "brand.toml").read_text(encoding="utf-8")


def test_theme_font_fallbacks():
    from brand_from_pptx import _fonts

    notes = []
    assert _fonts("+mn-lt", "Arial", notes) == {"heading": "Arial", "body": "Arial"}
    assert _fonts("Georgia", "", notes) == {"heading": "Georgia", "body": "Georgia"}
    neutral = load_brand(PLUGIN_BRANDS / "neutral").fonts["body"]
    assert _fonts("", "+mj-lt", notes) == {"heading": neutral, "body": neutral}
    assert len(notes) == 3


def test_draft_blocked_until_confirm(extracted, tmp_path):
    with pytest.raises(BrandError, match="draft"):
        load_brand(extracted.brand_dir)
    confirm_brand(extracted.brand_dir)
    toml = (extracted.brand_dir / "brand.toml").read_text(encoding="utf-8")
    assert "draft" not in toml.lower() and 'name = "cliente"' in toml
    brand = load_brand(extracted.brand_dir)
    assert brand.layouts["title"] == "Title Slide"
    assert not brand.draft
    d = Deck(brand)
    d.title_slide("k", "Título", "sub")
    s, top = d.content_slide("k", "Contenido")
    d.cards_row(s, top, [{"title": "A", "lines": ["uno"]}])
    assert d.save(tmp_path / "deck.pptx") == 2


def test_extract_refuses_existing_brand(extracted, client_pptx, tmp_path):
    with pytest.raises(BrandError, match="exists"):
        brand_from_pptx(client_pptx, "cliente", tmp_path / "brands", render_samples=False)


def test_master_images_copied(tmp_path, client_pptx):
    prs = Presentation(str(client_pptx))
    png = tmp_path / "logo.png"
    png.write_bytes(_PNG)
    slide = prs.slides[0]
    pic = slide.shapes.add_picture(str(png), Inches(0.2), Inches(0.2))
    rid = pic._element.blipFill.blip.rEmbed
    master = prs.slide_masters[0]
    new_rid = master.part.relate_to(slide.part.related_part(rid), RT.IMAGE)
    pic._element.blipFill.blip.set(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", new_rid)
    master.shapes._spTree.append(pic._element)
    src = tmp_path / "with-logo.pptx"
    prs.save(str(src))
    rep = brand_from_pptx(src, "logo", tmp_path / "brands", render_samples=False)
    assert [p.parent.name for p in rep.images] == ["assets-review"]
    assert rep.images[0].read_bytes() == _PNG


def test_extract_cli(client_pptx, tmp_path, capsys):
    import brand_from_pptx as bfp

    out = tmp_path / "cli-brands"
    assert bfp.main([str(client_pptx), "acme", "--out", str(out), "--no-samples"]) == 0
    printed = capsys.readouterr().out
    assert "layout 'Title Slide' / TextBox 3: 'Cliente X'" in printed and "sensitivity label: kept" in printed
    assert bfp.main(["--confirm", str(out / "acme")]) == 0
    assert not load_brand(out / "acme").draft
    assert bfp.main([str(client_pptx), "acme", "--out", str(out), "--no-samples"]) == 2  # exists


@pytest.mark.skipif(os.environ.get("OFFICE_DOCS_RENDER") != "1", reason="set OFFICE_DOCS_RENDER=1")
def test_extract_renders_samples(client_pptx, tmp_path):
    rep = brand_from_pptx(client_pptx, "cliente", tmp_path / "brands")
    assert rep.samples and all(p.suffix == ".png" and p.stat().st_size > 0 for p in rep.samples)


def test_new_brand_loads(tmp_path):
    folder = new_brand("acme", tmp_path, "#112233", "445566", "Arial")
    assert folder == tmp_path / "acme"
    brand = load_brand(folder)
    assert brand.colors["primary"] == "112233"
    assert brand.colors["accent"] == "445566"
    assert brand.fonts["body"] == "Arial" and brand.fonts["heading"] == "Arial"
    assert brand.style == "modern" and not brand.draft
    assert brand.colors["on_primary"] == "FFFFFF"


def test_new_brand_options(tmp_path):
    logo = tmp_path / "src-logo.png"
    logo.write_bytes(_PNG)
    folder = new_brand("light", tmp_path / "out", "F0F0F0", "#00aa77", "Segoe \"UI\" Light",
                       highlight="abcdef", heading_font="Georgia", logo=logo, style="outlined")
    brand = load_brand(folder)
    assert brand.fonts == {"body": 'Segoe "UI" Light', "heading": "Georgia", "mono": 'Segoe "UI" Light'}
    assert brand.colors["highlight"] == "ABCDEF"
    assert brand.colors["on_primary"] != "FFFFFF"  # a light primary needs dark text
    assert brand.colors["accent"] == brand.colors["card_border"] == "00AA77"  # outlined cards use the accent
    assert brand.logos["light"].parent == folder and brand.logos["light"].read_bytes() == _PNG


def test_new_brand_validates(tmp_path):
    with pytest.raises(BrandError, match="logo"):
        new_brand("x", tmp_path, "112233", "445566", "Arial", logo=tmp_path / "missing.png")
    with pytest.raises(BrandError, match="primary"):
        new_brand("x", tmp_path, "12345", "445566", "Arial")
    with pytest.raises(BrandError, match="style"):
        new_brand("x", tmp_path, "112233", "445566", "Arial", style="shiny")
    with pytest.raises(BrandError, match="name"):
        new_brand("../x", tmp_path, "112233", "445566", "Arial")
    assert not (tmp_path / "x").exists()
    new_brand("x", tmp_path, "112233", "445566", "Arial")
    with pytest.raises(BrandError, match="exists"):
        new_brand("x", tmp_path, "112233", "445566", "Arial")


def test_new_brand_deck_builds(tmp_path):
    d = Deck(new_brand("acme", tmp_path, "#112233", "445566", "Arial"))
    d.title_slide("Kicker", "Título", "Subtítulo")
    s, top = d.content_slide("K", "Cards")
    d.cards_row(s, top, [{"title": "Uno", "lines": ["a"]}, {"title": "Dos", "lines": ["b"]}])
    out = tmp_path / "acme.pptx"
    assert d.save(out) == 2
    assert len(Presentation(str(out)).slides) == 2


def test_new_brand_cli_samples(tmp_path, monkeypatch):
    import new_brand as nb

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(nb, "_sample_docx", lambda brand, folder: None)  # pandoc-free
    assert nb.main(["demo", "--primary", "112233", "--accent", "445566", "--font", "Arial",
                    "--out", str(tmp_path / "brands"), "--samples"]) == 0
    assert load_brand(tmp_path / "brands" / "demo").colors["primary"] == "112233"
    deck = tmp_path / "demo-samples" / "demo-sample.pptx"
    assert len(Presentation(str(deck)).slides) == 4
